import asyncio
import contextlib
import copy
import hashlib
import html
import io
import json
import mimetypes
import os
import secrets
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path
from typing import Literal

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from pydantic import BaseModel, Field, ConfigDict, StrictInt, StrictFloat, StrictStr, StrictBool

from . import domain, store, ax_service, drafting, rulebook, automation, download_names, accounts, notifications, intake_workflow, source_review

ROOT = store.ROOT
DEMO_MODE = os.getenv('DEBTOFF_DEMO_MODE', '1') == '1'
POOL = ThreadPoolExecutor(max_workers=1, thread_name_prefix='debtoff-agent')
COLLECTION_POOL = ThreadPoolExecutor(max_workers=1, thread_name_prefix='debtoff-collection')
LOGIN_ATTEMPTS = {}
LOGIN_LOCK = threading.Lock()
LEGAL_WATCH_LOCK = threading.Lock()
LEGAL_WATCH_JOB = None


def registry():
    return json.loads((ROOT / 'data/registry.json').read_text(encoding='utf-8'))


@asynccontextmanager
async def lifespan(app):
    store.initialize()
    accounts.bootstrap(DEMO_MODE)
    intake_workflow.migrate_stored_applications()
    if accounts.config(DEMO_MODE)['start_profile'] == 'demo':
        for case in domain.seed_cases(registry()):
            if not store.get_case(case['id']):
                store.insert_case(case)
    watcher = asyncio.create_task(_legal_watch_loop()) if os.getenv('DEBTOFF_LEGAL_WATCH', '0') == '1' else None
    yield
    if watcher:
        watcher.cancel()
        with contextlib.suppress(asyncio.CancelledError):
            await watcher


def _refresh_watch_cases():
    """Invalidate only applicable changed sources, then schedule affected cases in bounded batches."""
    from . import legal_watch
    with store.db() as con:
        rows = con.execute('SELECT body FROM cases').fetchall()
    for row in rows:
        case = json.loads(row['body'])
        policy_state = legal_watch.case_policy_status(case)
        signature = policy_state['signature']
        old_signature = case.get('legal_dependency_signature')
        if old_signature != signature:
            def apply(c):
                c['legal_dependency_signature'] = signature
                c['legal_update'] = policy_state
                if old_signature is not None or policy_state.get('pending_changes'):
                    domain.invalidate(c, '적용 법령·법원 기준 변경')
                    c['legal_change_pending'] = True
                    automation.notify(c, 'law:' + signature, '적용 근거 변경',
                        '사건에 적용되는 법령·법원 기준이 변경되어 서류 선정과 문서 근거를 다시 확인합니다.', kind='legal_update')
            try:
                case = store.mutate(case['id'], case['version'], ax_service.SYSTEM, 'law.version_checked', apply)
            except store.VersionConflict:
                continue
        if case.get('legal_change_pending'):
            with store.db() as con:
                queued = con.execute("SELECT count(*) FROM ax_runs WHERE status IN ('queued','running')").fetchone()[0]
            if queued >= 8:
                break
            ax_service.maybe_schedule(case, ax_service.SYSTEM, POOL, trigger='적용 법령·법원 기준 갱신')


async def _legal_watch_loop():
    from . import legal_watch
    while True:
        try:
            await legal_watch.run_once()
            await asyncio.to_thread(_refresh_watch_cases)
        except asyncio.CancelledError:
            raise
        except Exception:
            # Last-good source versions remain usable; the watcher's status retains
            # collection failures. A periodic retry must not terminate the API.
            pass
        await asyncio.sleep(60)


app = FastAPI(title='빚오프 공통 업무 API', version='0.2.0', lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=os.getenv('DEBTOFF_CORS_ORIGINS','http://localhost:5173,http://127.0.0.1:5173,http://localhost:5174,http://127.0.0.1:5174').split(','), allow_credentials=False, allow_methods=['GET','POST','PATCH','DELETE'], allow_headers=['Authorization','Content-Type'], expose_headers=['Content-Disposition'])


@app.middleware('http')
async def security_headers(request, call_next):
    # Demo credentials must never be exposed by an accidental public bind/deployment.
    if DEMO_MODE and request.client and request.client.host not in ('127.0.0.1','::1','testclient'):
        return JSONResponse({'detail':'Demo mode is available only from the local machine.'}, status_code=403)
    response = await call_next(request)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['X-Frame-Options'] = 'DENY'
    return response


@app.exception_handler(domain.DomainError)
async def domain_error(request, exc):
    return JSONResponse({'detail':exc.message,'code':exc.code}, status_code=422)


@app.exception_handler(store.VersionConflict)
async def conflict(request, exc):
    return JSONResponse({'detail':str(exc),'code':'VERSION_CONFLICT'}, status_code=409)


def current_user(authorization: str = Header(default='')):
    if not authorization.startswith('Bearer '):
        raise HTTPException(401, '로그인이 필요합니다.')
    token_hash = hashlib.sha256(authorization[7:].encode()).hexdigest()
    with store.db() as con:
        row = con.execute('SELECT * FROM sessions WHERE token_hash=? AND expires>?', (token_hash,time.time())).fetchone()
    if not row:
        raise HTTPException(401, '세션이 만료되었습니다. 다시 로그인하세요.')
    user = accounts.by_id(row['user_id'], DEMO_MODE)
    if not user:
        raise HTTPException(401, '계정이 회수되었습니다.')
    return user


def staff(user=Depends(current_user)):
    if user['role'] not in ('staff','lawyer'):
        raise HTTPException(403, '직원 권한이 필요합니다.')
    return user


def authorize(case_id, user):
    case = store.get_case(case_id)
    if not case or case['org_id'] != user['org_id']:
        raise HTTPException(404, '사건을 찾을 수 없습니다.')
    if user['role']=='client':
        allowed = case['client_user_id']==user['id']
    else:
        allowed = user['id'] in case['members']
    if not allowed:
        raise HTTPException(404, '사건을 찾을 수 없습니다.')
    return case


def visible(case, user):
    case = copy.deepcopy(case)
    if intake_workflow.pending(case):
        intake_workflow.waiting_state(case)
    case['notifications'] = notifications.projected(case, user)
    if user['role']=='client':
        fields = ('id','matter_number','title','client_name','court_id','court_name','stage','version','synthetic','assigned_to','requests','messages','documents')
        notices = case['notifications']
        intake = {k: v for k, v in case.get('intake', {}).items()
                  if k in ('status', 'submitted_at', 'requested_at', 'completed_at')}
        case = {k:case[k] for k in fields if k in case}
        if intake:
            case['intake'] = intake
        case['notifications'] = notices
        request_fields = ('id','version','catalog_id','title','issuer','institution','account_masked','period',
                          'period_start','period_end','issuance_options','status','issuance_url','steps',
                          'not_proven','document_ids','due_date','public_review_note','withdrawal_reason','no_longer_required',
                          'upload_mode','supports_multiple_files','upload_examples','collection_progress','requirement_kind','qualification')
        case['requests'] = [{k:r.get(k) for k in request_fields if k in r} for r in case['requests']]
        case['documents']=[{k:d.get(k) for k in ('id','filename','request_id','status','created_at','public_status')} for d in case['documents'] if d['uploader']==user['id']]
        allowed_doc_ids={d['id'] for d in case['documents']}
        for r in case['requests']:
            r['document_ids']=[d for d in r['document_ids'] if d in allowed_doc_ids]
        return case
    case['blockers']=domain.blockers(case)
    case['deadline_alerts']=domain.deadline_alerts(case)
    from .approval_estimator import estimate
    case['approval_estimate'] = estimate(case)
    case['ax_runs']=ax_service.case_runs(case)
    ax_service.project_active_pipeline(case, case['ax_runs'])
    from . import extraction_readiness
    extraction_readiness.project(case)
    if not case.get('ax_pipeline'):
        # Older cases still show their actual waiting stage before their first
        # AX event. Reading a case never fabricates a completed verification.
        preview = copy.deepcopy(case)
        requests = [r for r in case.get('requests', []) if r.get('status') not in automation.INACTIVE and not r.get('no_longer_required')]
        awaiting_upload = not requests or any(r.get('status') in ('requested', 'needs_more') for r in requests)
        case['ax_pipeline'] = automation._state(preview, 'collecting' if awaiting_upload else 'verification_waiting', 0 if awaiting_upload else 1)
    from . import workflow_contract
    definitions = {step['id']: step for step in workflow_contract.STAGES}
    # Enrich the read projection with present-day explanations, preserving the
    # recorded execution version/status of historical runs.
    pipeline = case.get('ax_pipeline', {})
    pipeline['display_contract_version'] = workflow_contract.VERSION
    for step in pipeline.get('steps', []):
        definition = definitions.get(step.get('id'), {})
        for field in ('inputs', 'outputs', 'gate', 'on_failure'):
            if field not in step and field in definition:
                step[field] = copy.deepcopy(definition[field])
    case['rule_evaluation']=rulebook.evaluate_case(case)
    from . import dashboard_insights
    case['case_assessment'] = dashboard_insights.build(case)
    for field in ('extraction_candidates','drafts','automation_history','tasks','legal_calculations','court_documents'):
        case.setdefault(field,[])
    case.setdefault('automation',{})
    with store.db() as con:
        case['agent_runs']=[present_run(r,case) for r in con.execute('SELECT body,version FROM runs WHERE case_id=? ORDER BY rowid DESC LIMIT 20',(case['id'],))]
    return case


def present_run(row,case):
    run=json.loads(row['body'])
    run['case_version']=row['version']
    if row['version']!=case['version'] and run['status'] not in ('queued','running','failed'):
        run.update(status='stale',error='실행 이후 사건이 변경되었습니다. 현재 근거로 다시 검토하세요.')
    return run


def change(case_id, user, version, action, fn):
    authorize(case_id,user)
    def apply(case):
        fn(case)
        notifications.record_event(case, user, action)
        if intake_workflow.pending(case):
            intake_workflow.waiting_state(case)
    saved=store.mutate(case_id, version, user, action, apply)
    if action in ('document.received','document.verified','document.metadata','document.scopes','document.request_assigned','consultation.recorded','correction.created','request.created','request.withdrawn','fact.confirmed','issue.decided','legal_calculation.created','candidate.reviewed','intake.reviewed') or (action=='message.created' and user['role']=='client'):
        ax_service.maybe_schedule(saved,user,POOL,trigger={'document.received':'새 서류 제출','document.verified':'서류 확인 결과 반영','consultation.recorded':'상담 기록 변경','correction.created':'보정 요구 등록','message.created':'내담자 회신'}.get(action, '보완 자료·판단 반영'))
    return visible(saved,user)


class Payload(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expected_version: StrictInt = Field(ge=1)


class Login(BaseModel):
    username: str = Field(max_length=100)
    password: str = Field(max_length=200)


@app.post('/api/auth/login')
def login(data: Login, request: Request):
    key = request.client.host if request.client else 'unknown'
    with LOGIN_LOCK:
        attempts = [t for t in LOGIN_ATTEMPTS.get(key,[]) if time.time()-t < 60]
        LOGIN_ATTEMPTS[key]=attempts
        if len(attempts)>=20:
            raise HTTPException(429,'로그인 시도가 많습니다. 잠시 후 다시 시도하세요.')
        attempts.append(time.time())
    user = accounts.authenticate(data.username, data.password, DEMO_MODE)
    if not user:
        raise HTTPException(401,'계정 또는 비밀번호를 확인하세요.')
    token = secrets.token_urlsafe(40)
    expires = time.time()+8*3600
    with store.db() as con:
        con.execute('DELETE FROM sessions WHERE expires<?',(time.time(),))
        con.execute('INSERT INTO sessions VALUES (?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),user['id'],expires))
    return {'token':token,'user':user,'expires_at':expires}


@app.get('/api/auth/config')
def auth_config():
    return accounts.config(DEMO_MODE)


class AccountCreate(BaseModel):
    model_config = ConfigDict(extra='forbid')
    username: str = Field(min_length=3, max_length=80)
    name: str = Field(min_length=1, max_length=80)
    role: Literal['client', 'staff', 'lawyer'] = 'client'
    password: str = Field(min_length=16, max_length=200)


@app.post('/api/auth/accounts')
def create_account(data: AccountCreate, user=Depends(staff)):
    return accounts.create(user, **data.model_dump())


@app.get('/api/auth/me')
def me(user=Depends(current_user)):
    return user


@app.post('/api/auth/logout')
def logout(authorization: str=Header(default=''),user=Depends(current_user)):
    with store.db() as con:
        con.execute('DELETE FROM sessions WHERE token_hash=?',(hashlib.sha256(authorization[7:].encode()).hexdigest(),))
    return {'ok':True}


@app.get('/api/health')
def health():
    return {'status':'ok','service':'debtoff',**accounts.config(DEMO_MODE),'version':'0.2.0','legal_rules_active':False}


@app.get('/api/notifications')
def notification_feed(limit: int = 100, unread_only: bool = False, user=Depends(current_user)):
    domain.require(1 <= limit <= 100, 'PAGINATION', '알림은 한 번에 1~100건 조회합니다.')
    with store.db() as con:
        rows = con.execute('SELECT body FROM cases WHERE org_id=?', (user['org_id'],)).fetchall()
    notices = []
    for row in rows:
        case = json.loads(row['body'])
        if not ((user['role'] == 'client' and case.get('client_user_id') == user['id'])
                or (user['role'] != 'client' and user['id'] in case.get('members', []))):
            continue
        notices.extend({**notice, 'case_id': case['id'], 'case_title': case['title'],
                        'matter_number': case.get('matter_number'), 'case_version': case['version']}
                       for notice in notifications.projected(case, user))
    notices.sort(key=lambda row: row.get('created_at', ''), reverse=True)
    unread_count = sum(row['unread'] for row in notices)
    filtered = [row for row in notices if row['unread']] if unread_only else notices
    return {'notifications': filtered[:limit], 'unread_count': unread_count, 'total': len(filtered)}


@app.get('/api/cases')
def cases(limit:int=50,offset:int=0,user=Depends(current_user)):
    domain.require(1<=limit<=100 and offset>=0,'PAGINATION','목록은 한 번에 1~100건 조회합니다.')
    with store.db() as con:
        rows=con.execute('SELECT body FROM cases WHERE org_id=?',(user['org_id'],)).fetchall()
    result=[]
    for row in rows:
        case=json.loads(row['body'])
        if (user['role']=='client' and case['client_user_id']==user['id']) or (user['role']!='client' and user['id'] in case['members']):
            summary={k:case[k] for k in ('id','title','client_name','court_id','court_name','stage','version','synthetic','assigned_to','updated_at','summary')}
            summary['matter_number'] = case.get('matter_number')
            notice_rows = notifications.projected(case, user)
            summary['unread_count'] = sum(n['unread'] for n in notice_rows)
            blocked=domain.blockers(case)
            summary.update(blocker_count=len(blocked),next_action=blocked[0]['title'] if blocked else '검토 보고서 준비',correction_count=sum(x['status']!='approved' for x in case['corrections']),document_count=len(case['documents']),request_count=len(case['requests']),deadline_alerts=domain.deadline_alerts(case))
            if user['role']=='client':
                for internal in ('summary','deadline_alerts','next_action','blocker_count','correction_count'):
                    summary.pop(internal,None)
                summary['document_count']=sum(d['uploader']==user['id'] for d in case['documents'])
            else:
                summary['review_count'] = sum(n['unread'] and n.get('kind') in {'review_request', 'customer_handoff'} for n in notice_rows)
                runs=ax_service.case_runs(case,1)
                latest=runs[0] if runs else {}
                summary.update(ax_status=latest.get('status') or case.get('automation',{}).get('stage'),ax_finding_count=len(latest.get('findings',[])),ax_last_at=latest.get('finished_at') or latest.get('created_at'),case_type=case.get('case_type'))
                if intake_workflow.pending(case):
                    # Historical premature analysis remains in the case history;
                    # the dashboard's next action follows the current interview.
                    summary.update(ax_status=None, ax_finding_count=0, ax_last_at=None,
                                   next_action='세부 상담 요청' if case['intake']['status'] == 'received' else '세부 상담 기록')
            result.append(summary)
    return {'cases':result[offset:offset+limit],'total':len(result),'limit':limit,'offset':offset}


@app.get('/api/cases/{case_id}')
def case_detail(case_id: str,user=Depends(current_user)):
    case=authorize(case_id,user)
    store.access(user,case_id,'case.read')
    return visible(case,user)


class CaseCreate(BaseModel):
    client_name: str=Field(min_length=1,max_length=60)
    court_id: str
    summary: str=Field(default='',max_length=3000)
    consent: bool


@app.post('/api/cases')
def create_case(data:CaseCreate,user=Depends(staff)):
    domain.require(data.consent,'CONSENT_REQUIRED','동의 확인이 필요합니다.')
    court=next((x for x in registry()['courts'] if x['id']==data.court_id),None)
    domain.require(court is not None,'COURT_UNKNOWN','관할 후보를 선택하세요.')
    case=domain.new_case(data.client_name,data.court_id,court['name'],data.summary,DEMO_MODE)
    case['consent']={'confirmed_by':user['id'],'at':store.now(),'version':'local-v1'}
    store.insert_case(case)
    return visible(case,user)


@app.get('/api/registry')
def get_registry(user=Depends(staff)):
    return registry()


@app.get('/api/requirements')
def requirements(user=Depends(staff)):
    return json.loads((ROOT/'data/requirements.json').read_text(encoding='utf-8'))


class RequestDocument(Payload):
    catalog_id: str
    period: str=Field(min_length=1,max_length=500)
    reason: str=Field(min_length=1,max_length=1000)


class Consultation(Payload):
    notes: str=Field(default='',max_length=10000)
    answers: dict[str,str]
    consent: bool


class ConsultationRequest(Payload):
    message: str = Field(min_length=10, max_length=4000)


@app.post('/api/cases/{case_id}/consultation/request')
def request_consultation(case_id: str, data: ConsultationRequest, user=Depends(staff)):
    return change(case_id, user, data.expected_version, 'consultation.requested',
                  lambda c: intake_workflow.request_consultation(c, user, data.message))


@app.post('/api/cases/{case_id}/consultation')
def consultation(case_id:str,data:Consultation,user=Depends(staff)):
    allowed={'income','household','assets','debts','housing','recent_transactions','history','unknowns'}
    domain.require(set(data.answers)<=allowed and all(len(v)<=3000 for v in data.answers.values()),'QUESTION_SCOPE','질문 항목과 입력 길이를 확인하세요.')
    domain.require(data.consent,'CONSENT_REQUIRED','상담 기록 동의 확인이 필요합니다.')
    domain.require(len('\n'.join([data.notes.strip(), *[v.strip() for v in data.answers.values()]]).strip()) >= 10,
                   'CONSULTATION_CONTENT', '세부 상담에서 확인한 내용을 10자 이상 기록하세요.')
    def apply(c):
        intake_workflow.complete(c, user)
        domain.invalidate(c,'상담 진술 변경')
        c['consultation']={'notes':data.notes,'answers':data.answers,'consent':data.consent,'recorded_by':user['name'],'updated_at':store.now(),'status':'party_statement','version':c.get('consultation',{}).get('version',0)+1}
    return change(case_id,user,data.expected_version,'consultation.recorded',apply)


class Workflow(Payload):
    stage: Literal['자료 수집','검증 중','보정 대응','검토 보류']
    reason: str=Field(min_length=5,max_length=1000)


@app.post('/api/cases/{case_id}/workflow')
def workflow(case_id:str,data:Workflow,user=Depends(staff)):
    def apply(c):
        c.update(stage=data.stage,stage_reason=data.reason)
    return change(case_id,user,data.expected_version,'workflow.changed',apply)


@app.post('/api/cases/{case_id}/requests')
def request_document(case_id:str,data:RequestDocument,user=Depends(staff)):
    doc=next((d for d in registry()['documents'] if d['id']==data.catalog_id),None)
    domain.require(doc is not None,'CATALOG_UNKNOWN','문서 종류를 선택하세요.')
    def apply(case):
        domain.invalidate(case,'자료 요구 변경')
        domain.request_document(case,doc,data.period,data.reason)
        automation.sync_request_notifications(case)
    return change(case_id,user,data.expected_version,'request.created',apply)


def extract_file(content, extension):
    pages=[]
    try:
        if extension=='.pdf':
            domain.require(content.startswith(b'%PDF-'),'FILE_FORMAT','PDF 형식이 올바르지 않습니다.')
            import fitz
            with fitz.open(stream=content,filetype='pdf') as pdf:
                domain.require(not pdf.is_encrypted,'FILE_ENCRYPTED','암호화된 PDF는 해제 후 다시 제출하세요.')
                domain.require(len(pdf)<=100,'TOO_MANY_PAGES','한 파일은 100쪽 이하로 제출하세요.')
            from .ocr_engine import extract_pages
            pages=extract_pages(content,extension)
        elif extension in ('.txt','.csv'):
            try:
                text=content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text=content.decode('cp949')
            pages=[{'page':1,'text':text}]
        elif extension=='.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                domain.require(sum(i.file_size for i in z.infolist())<30_000_000,'ARCHIVE_LIMIT','압축 해제 크기가 너무 큽니다.')
                domain.require(not any('vba' in i.filename.lower() for i in z.infolist()),'MACRO_BLOCKED','매크로 문서는 받을 수 없습니다.')
            from docx import Document
            doc=Document(io.BytesIO(content))
            text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
            pages=[{'page':1,'text':text}]
        elif extension in ('.png','.jpg','.jpeg'):
            domain.require(content.startswith(b'\x89PNG\r\n\x1a\n') if extension=='.png' else content.startswith(b'\xff\xd8\xff'),'FILE_FORMAT','이미지 형식이 올바르지 않습니다.')
            from .ocr_engine import extract_pages
            pages=extract_pages(content,extension)
        else:
            raise domain.DomainError('FILE_TYPE','PDF, DOCX, TXT, CSV, PNG, JPG 파일을 지원합니다.')
    except domain.DomainError:
        raise
    except Exception:
        raise domain.DomainError('FILE_PARSE_FAILED','파일을 읽지 못했습니다. 원문 형식·암호를 확인하거나 읽을 수 있는 파일로 제출하세요.')
    return pages


@app.post('/api/cases/{case_id}/documents')
async def upload(case_id:str,file:UploadFile=File(...),request_id:str=Form(default=''),expected_version:int=Form(...),replaces_document_id:str=Form(default=''),user=Depends(current_user)):
    return await _receive_documents(case_id,[file],request_id,expected_version,user,replaces_document_id)


@app.post('/api/cases/{case_id}/documents/batch')
async def upload_batch(case_id:str,files:list[UploadFile]=File(...),request_id:str=Form(default=''),expected_version:int=Form(...),replaces_document_id:str=Form(default=''),user=Depends(current_user)):
    return await _receive_documents(case_id,files,request_id,expected_version,user,replaces_document_id)


async def _receive_documents(case_id,files,request_id,expected_version,user,replaces_document_id=''):
    """Parse a bounded batch, then commit all originals in one case transaction.

    Additional files supplement a request. Replacing an earlier original is an
    explicit operation; adding December's payroll must never hide November's.
    A failed parse, duplicate or version conflict stores none of this batch.
    """
    case=authorize(case_id,user)
    if case['version']!=expected_version:
        raise store.VersionConflict('사건이 변경되었습니다. 새로고침 후 제출하세요.')
    domain.require(1<=len(files)<=20,'FILE_BATCH_SIZE','한 번에 1~20개 파일을 제출할 수 있습니다.')
    domain.require(not replaces_document_id or len(files)==1,'REPLACEMENT_SINGLE_FILE','원본 교체는 한 번에 한 파일씩 지정해 주세요.')
    prepared=[]
    total_bytes=0
    for file in files:
        content=await file.read(10_000_001)
        domain.require(0<len(content)<=10_000_000,'FILE_SIZE','각 파일은 0바이트 초과 10MB 이하여야 합니다.')
        total_bytes+=len(content)
        domain.require(total_bytes<=50_000_000,'FILE_BATCH_BYTES','한 번에 제출하는 파일의 합계는 50MB 이하여야 합니다.')
        filename=Path((file.filename or 'document').replace('\\','/')).name[:160]
        extension=Path(filename).suffix.lower()
        pages=await asyncio.to_thread(extract_file,content,extension)
        prepared.append({'id':store.uid('doc'),'filename':filename,'extension':extension,
            'content':content,'pages':pages,'sha256':hashlib.sha256(content).hexdigest(),'request_id':request_id})
    return _commit_prepared_documents(case_id,prepared,expected_version,user,replaces_document_id)


def _commit_prepared_documents(case_id,prepared,expected_version,user,replaces_document_id=''):
    """Shared atomic commit for selected-file batches and content-routed imports."""
    case=authorize(case_id,user)
    if case['version']!=expected_version:
        raise store.VersionConflict('사건이 변경되었습니다. 새로고침 후 제출하세요.')
    domain.require(bool(prepared),'NO_IMPORT_DOCUMENTS','제출할 수 있는 서류가 없습니다.')
    domain.require(len({row['sha256'] for row in prepared})==len(prepared),'DUPLICATE_FILE','선택한 파일 중 내용이 같은 파일이 있습니다. 중복 파일을 제외해 주세요.')
    folder=store.DATA_DIR/'uploads'/case_id
    folder.mkdir(parents=True,exist_ok=True)
    targets=[]
    def apply(c):
        identities={(row['sha256'],row.get('request_id','')) for row in prepared}
        domain.require(not any((d.get('sha256'),d.get('request_id','')) in identities
            and d.get('status') not in {'superseded','rejected'} for d in c['documents']),
            'DUPLICATE_FILE','같은 요청에 동일한 파일이 이미 등록되어 있습니다.')
        requests={}
        for row in prepared:
            request_id=row.get('request_id','')
            request_item=domain.item(c,'requests',request_id) if request_id else None
            domain.require(not request_item or (request_item['status'] not in automation.INACTIVE and not request_item.get('no_longer_required')),
                           'REQUEST_WITHDRAWN', '철회된 요청에는 제출할 수 없습니다. 현재 요청을 선택하세요.')
            if request_item:requests[request_id]=request_item
        if replaces_document_id:
            domain.require(len(prepared)==1,'REPLACEMENT_SINGLE_FILE','원본 교체는 한 번에 한 파일씩 지정해 주세요.')
            request_id=prepared[0].get('request_id','')
            previous=domain.item(c,'documents',replaces_document_id)
            domain.require(previous.get('request_id','')==request_id,'REPLACEMENT_REQUEST_MISMATCH','같은 요청에 제출한 원본만 교체할 수 있습니다.')
            domain.require(previous.get('status') not in {'superseded','rejected','quarantined'},'INACTIVE_DOCUMENT','현재 사용 중인 원본을 선택해 주세요.')
            if user['role']=='client':
                domain.require(previous.get('uploader')==user['id'],'REPLACEMENT_OWNER','본인이 제출한 원본만 교체할 수 있습니다.')
            previous.setdefault('status_history',[]).append({'at':store.now(),'status':previous.get('status'),'reason':'사용자가 원본 교체를 명시적으로 요청'})
            previous.update(status='superseded',superseded_by=prepared[0]['id'])
        domain.invalidate(c,'자료 수신으로 근거 버전 변경')
        from . import extraction_readiness
        for row in prepared:
            pages=row['pages']
            document={'id':row['id'],'filename':row['filename'],'storage_name':row['id']+row['extension'],
                'sha256':row['sha256'],'size':len(row['content']),'version':1,'request_id':row.get('request_id',''),
                'status':'received','public_status':'파일 수신 · 담당자 확인 전','created_at':store.now(),
                'uploader':user['id'],'text':'\n'.join(p['text'] for p in pages),'page_texts':pages,
                'extraction_status':'extracted' if any(p['text'].strip() for p in pages) else 'manual_review',
                'scope_confirmed':False,'content_confirmed':False,'person_confirmed':False,
                'security_status':'extension_and_signature_checked; antivirus_not_configured'}
            if row.get('import_routing'):
                document['import_routing']=copy.deepcopy(row['import_routing'])
                document['classification_review_required']=row['import_routing'].get('status')!='matched'
                document['scope_review_required']=bool(row['import_routing'].get('scope_pending'))
                if document['classification_review_required']:
                    document['public_status']='파일 수신 · 연결할 서류 항목 확인 필요'
                if row['import_routing'].get('identity_status')=='conflict':
                    document.update(status='quarantined',public_status='명의 확인 후 연결 필요')
                    document['automated_check']={'coverage_status':'identity_conflict'}
            if row.get('import_job_id'):document['import_job_id']=row['import_job_id']
            if replaces_document_id:document['replaces_document_id']=replaces_document_id
            methods={p.get('extraction_method','native_text') for p in pages}
            pending=[p['page'] for p in pages if not p.get('text','').strip()]
            document.update(extraction_method='mixed' if len(methods)>1 else next(iter(methods),'unreadable'),
                ocr_summary={'ocr_pages':[p['page'] for p in pages if p.get('extraction_method')=='ocr'],
                'unreadable_pages':pending,'page_count':len(pages),'review_required':any(p.get('requires_review') for p in pages)})
            if pending and any(p.get('text','').strip() for p in pages):document['extraction_status']='partial'
            c['documents'].append(document)
            extracted=extraction_readiness.candidates(document)
            ax_service.enrich_candidates(c,extracted)
            extraction_readiness.capture(c,document,extracted)
        for request_id,request_item in requests.items():
            request_item.setdefault('document_ids',[]).extend(row['id'] for row in prepared if row.get('request_id','')==request_id)
            request_item['status']='received'
            request_item.pop('automatic_validation',None)
            from .request_collection import sync_review_status
            sync_review_status(c,request_item)
    try:
        for row in prepared:
            target=(folder/(row['id']+row['extension'])).resolve()
            domain.require(target.is_relative_to(folder.resolve()),'ARTIFACT_PATH','잘못된 저장 경로입니다.')
            targets.append(target)
            target.write_bytes(row['content'])
        return change(case_id,user,expected_version,'document.received',apply)
    except Exception:
        # change() commits before scheduling and building the response. A
        # post-commit display/worker error must never delete stored evidence.
        # If persistence cannot be inspected, retaining an orphan is safer
        # than removing a possibly committed original.
        try:
            saved=store.get_case(case_id)
            committed={doc.get('id') for doc in (saved or {}).get('documents',[])}
        except Exception:
            committed={row['id'] for row in prepared}
        for row,target in zip(prepared,targets):
            if row['id'] not in committed:target.unlink(missing_ok=True)
        raise


@app.get('/api/cases/{case_id}/documents/{doc_id}/download')
def download(case_id:str,doc_id:str,user=Depends(current_user)):
    case=authorize(case_id,user)
    doc=domain.item(case,'documents',doc_id)
    # Clients may only download their own uploaded originals.
    if user['role']=='client' and doc['uploader']!=user['id']:
        raise HTTPException(403,'이 원문은 내부 검토 자료입니다.')
    store.access(user,case_id,'document.download:'+doc_id)
    return FileResponse(store.DATA_DIR/'uploads'/case_id/doc['storage_name'],filename=doc['filename'],media_type='application/octet-stream')


class Verify(Payload):
    scope_confirmed: bool
    content_confirmed: bool
    person_confirmed: bool
    reason: str=Field(min_length=5,max_length=1500)


@app.post('/api/cases/{case_id}/documents/{doc_id}/verify')
def verify(case_id:str,doc_id:str,data:Verify,user=Depends(staff)):
    def apply(c):
        doc=source_review.reviewable_document(c,doc_id)
        domain.invalidate(c,'원문 검증 변경')
        source_review.apply_document(c,doc,data,user)
    return change(case_id,user,data.expected_version,'document.verified',apply)


class FactConfirm(Payload):
    value: StrictInt | None
    evidence_ids: list[str]
    reason: str=Field(min_length=1,max_length=1500)


@app.post('/api/cases/{case_id}/facts/{fact_id}/confirm')
def fact_confirm(case_id:str,fact_id:str,data:FactConfirm,user=Depends(staff)):
    return change(case_id,user,data.expected_version,'fact.confirmed',lambda c:domain.confirm_fact(c,fact_id,data.model_dump()))


class Decision(Payload):
    decision: str=Field(min_length=1,max_length=1500)
    reason: str=Field(min_length=5,max_length=3000)
    evidence_ids: list[str]


@app.post('/api/cases/{case_id}/issues/{issue_id}/decide')
def issue_decide(case_id:str,issue_id:str,data:Decision,user=Depends(staff)):
    return change(case_id,user,data.expected_version,'issue.decided',lambda c:domain.decide(c,issue_id,data.model_dump(),user))


class Message(Payload):
    text: str=Field(min_length=1,max_length=4000)


@app.post('/api/cases/{case_id}/messages')
def message(case_id:str,data:Message,user=Depends(current_user)):
    def apply(c):
        c['messages'].append({'id':store.uid('msg'),'text':data.text,'sender':user['name'],'role':user['role'],'created_at':store.now()})
        if user['role']=='client':
            domain.invalidate(c,'새 고객 진술 수신 · 영향 확인 필요')
            if not any(i['id']=='issue-client-reply' for i in c['issues']):
                c['issues'].append({'id':'issue-client-reply','title':'새 고객 회신의 사실·법률 영향 확인','description':'새로운 고객 진술을 확인하고 소득·채무·생활 상황 및 기존 판단에 영향이 있는지 검토하세요.','status':'open','evidence_ids':[],'counter_evidence':['고객 진술은 원문 증빙 확인 전'],'options':['원문 및 영향 확인','추가 자료 요청'],'decision':None,'reason':''})
        if user['role']=='client' and any(k in data.text for k in ('완납','가족','갚았')):
            if not any(i['id']=='issue-reply' for i in c['issues']):
                domain.invalidate(c,'새 고객 진술의 법률 검토 필요')
                c['issues'].append({'id':'issue-reply','title':'상환·가족지원 회신 확인','description':'고객 회신을 확인하고 완납 여부·지급자·자금출처·가족간 법률관계에 맞게 자료 요구를 검토하세요.','status':'open','evidence_ids':[],'counter_evidence':['원문 증빙 확인 전 고객 진술'],'options':['완납증명·송금증빙 요청','채권 변동과 대위 관계 검토'],'decision':None,'reason':''})
    return change(case_id,user,data.expected_version,'message.created',apply)


class Calculation(Payload):
    mode: Literal['scenario','operational']='scenario'
    income: StrictInt
    expenses: StrictInt
    months: StrictInt
    liquidation_value: StrictInt


@app.post('/api/cases/{case_id}/calculate')
def calculate(case_id:str,data:Calculation,user=Depends(staff)):
    return change(case_id,user,data.expected_version,'calculation.created',lambda c:domain.calculate(c,data.model_dump()))


@app.post('/api/cases/{case_id}/bundles')
def bundles(case_id:str,data:Payload,user=Depends(staff)):
    return change(case_id,user,data.expected_version,'bundle.created',domain.make_bundle)


class Reason(Payload):
    reason: str=Field(min_length=5,max_length=3000)


@app.post('/api/cases/{case_id}/bundles/{bundle_id}/approve')
def bundle_approve(case_id:str,bundle_id:str,data:Reason,user=Depends(staff)):
    return change(case_id,user,data.expected_version,'bundle.review_approved',lambda c:domain.approve_bundle(c,bundle_id,data.model_dump(),user))


@app.get('/api/cases/{case_id}/bundles/{bundle_id}/download')
def bundle_download(case_id:str,bundle_id:str,user=Depends(staff)):
    case=authorize(case_id,user)
    bundle=domain.item(case,'bundles',bundle_id)
    store.access(user,case_id,'bundle.download:'+bundle_id)
    e=lambda value:html.escape(str(value))
    facts=''.join(f"<tr><td>{e(f['label'])}</td><td>{e(f['value'] if f['value'] is not None else '미확인')}</td><td>{e(f['status'])}</td><td>{e(', '.join(f['evidence_ids']))}</td></tr>" for f in bundle['snapshot']['facts'])
    issues=''.join(f"<article><h3>{e(i['title'])}</h3><p>{e(i['description'])}</p><p>판단: {e(i.get('decision') or '검토 전')} / 사유: {e(i.get('reason') or '미기록')}</p></article>" for i in bundle['snapshot']['issues'])
    corrections=''.join(f"<tr><td>{e(i['title'])}</td><td>{e(i['requirement'])}</td><td>{e(i['status'])}</td><td>{e(i.get('answer') or '미작성')}</td><td>{e(', '.join(i.get('evidence_ids',[])))}</td></tr>" for i in bundle['snapshot']['corrections'])
    body=f'<!doctype html><html lang="ko"><meta charset="utf-8"><title>빚오프 검토 보고서</title><style>body{{font:15px sans-serif;max-width:960px;margin:40px auto;line-height:1.7}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ddd;padding:9px}}.note{{background:#fff3dc;padding:16px}}@media print{{body{{margin:0}}}}</style><h1>사건 검증 보고서</h1><p>{e(case["title"])} · {e(case["court_name"])}</p><p class="note">검토용 초안 · 법원 제출서식 아님. {"입력 변경으로 만료됨" if bundle["stale"] else e(bundle["status"])}</p><p>입력 리비전 {bundle["input_revision"]} · SHA-256 {bundle["snapshot_hash"]}</p><h2>확인 사실과 원문 근거</h2><table><tr><th>항목</th><th>값</th><th>상태</th><th>문서 ID</th></tr>{facts}</table><h2>법률 쟁점 및 반대 자료</h2>{issues}<h2>보정 요구 충족표</h2><table><tr><th>항목</th><th>요구</th><th>상태</th><th>답변</th><th>증거</th></tr>{corrections}</table><p>근거 없는 사실·법률 결론은 추가하지 않았습니다. 최종 서식·계산·관할 기준은 담당 변호사 검토가 필요합니다.</p></html>'
    return HTMLResponse(body,headers={'Content-Disposition':f'attachment; filename="{bundle_id}.html"','Content-Security-Policy':"default-src 'none'; style-src 'unsafe-inline'"})


class Approval(Reason):
    gate: Literal['H1','H-R']


@app.post('/api/cases/{case_id}/approvals')
def approval(case_id:str,data:Approval,user=Depends(staff)):
    domain.require(user['role']=='lawyer','LAWYER_ONLY','담당 변호사의 승인이 필요합니다.')
    def apply(c):
        c['approvals'].append({'id':store.uid('approval'),'gate':data.gate,'input_revision':c['input_revision'],'reason':data.reason,'actor':user['name'],'at':store.now(),'stale':False})
    return change(case_id,user,data.expected_version,'approval.created',apply)


class Correction(Payload):
    title: str=Field(min_length=1,max_length=200)
    source_document_id: str
    source_page: int=Field(ge=1,le=100)
    requirement: str=Field(min_length=5,max_length=4000)
    period: str=Field(min_length=1,max_length=500)
    due_date: date | None=None
    deadline_evidence_id: str | None=None


@app.post('/api/cases/{case_id}/corrections')
def correction(case_id:str,data:Correction,user=Depends(staff)):
    def apply(c):
        source=domain.evidence(c,[data.source_document_id],verified=False)[0]
        domain.require(not source['page_texts'] or data.source_page<=len(source['page_texts']),'PAGE_UNKNOWN','원문 페이지를 확인하세요.')
        if data.due_date:
            domain.evidence(c,[data.deadline_evidence_id] if data.deadline_evidence_id else [])
        domain.invalidate(c,'보정 요구 등록')
        record=data.model_dump(mode='json',exclude={'expected_version'})
        deadline_id=store.uid('deadline')
        record.update(id=store.uid('corr'),round=1,status='open',answer='',evidence_ids=[],deadline_id=deadline_id)
        c['corrections'].append(record)
        c['deadlines'].append({'id':deadline_id,'title':data.title,'kind':'legal','due_date':data.due_date.isoformat() if data.due_date else None,'evidence_id':data.deadline_evidence_id,'status':'confirmed' if data.due_date else 'unconfirmed','assigned_to':user['name'],'backup':'담당 변호사'})
    return change(case_id,user,data.expected_version,'correction.created',apply)


class CorrectionResponse(Reason):
    answer: str=Field(min_length=5,max_length=5000)
    evidence_ids: list[str]


@app.post('/api/cases/{case_id}/corrections/{correction_id}/respond')
def correction_respond(case_id:str,correction_id:str,data:CorrectionResponse,user=Depends(staff)):
    def apply(c):
        corr=domain.item(c,'corrections',correction_id)
        domain.evidence(c,data.evidence_ids)
        domain.require(corr.get('source_document_id'),'SOURCE_REQUIRED','실제 보정 원문을 연결한 항목만 충족 검토할 수 있습니다. 합성 예시 항목은 참고용입니다.')
        # Replies are independent work products; editing B must not make A impossible to complete.
        for bundle in c['bundles']:
            bundle.update(stale=True,stale_reason='보정 답변 변경')
        for approval in c['approvals']:
            if approval.get('correction_id')==correction_id or approval.get('bundle_id'):
                approval.update(stale=True,stale_reason='보정 답변 변경')
        corr.update(answer=data.answer,evidence_ids=data.evidence_ids,reason=data.reason,status='reviewed',input_revision=c['input_revision'],response_revision=corr.get('response_revision',0)+1)
    return change(case_id,user,data.expected_version,'correction.reviewed',apply)


@app.post('/api/cases/{case_id}/corrections/{correction_id}/approve')
def correction_approve(case_id:str,correction_id:str,data:Reason,user=Depends(staff)):
    domain.require(user['role']=='lawyer','LAWYER_ONLY','담당 변호사만 보정안을 승인할 수 있습니다.')
    def apply(c):
        corr=domain.item(c,'corrections',correction_id)
        domain.require(corr['status']=='reviewed' and corr.get('input_revision')==c['input_revision'],'CORRECTION_NOT_READY','원문 요구와 현재 증거 버전의 충족 검토가 필요합니다.')
        domain.require(corr.get('due_date'),'DEADLINE_UNCONFIRMED','실제 송달 근거와 기한을 확인하세요.')
        domain.evidence(c,[corr.get('source_document_id')])
        domain.evidence(c,corr.get('evidence_ids',[]))
        domain.evidence(c,[corr.get('deadline_evidence_id')])
        domain.require(not any(i['status']!='decided' for i in c['issues']),'LEGAL_ISSUE_OPEN','미결 법률 쟁점부터 판단하세요.')
        corr.update(status='approved',approved_by=user['name'],approval_reason=data.reason,approved_at=store.now())
        c['approvals'].append({'id':store.uid('approval'),'gate':'H-C','correction_id':correction_id,'input_revision':c['input_revision'],'reason':data.reason,'actor':user['name'],'at':store.now(),'stale':False})
    return change(case_id,user,data.expected_version,'correction.approved',apply)


class DeadlineConfirm(Reason):
    due_date: date
    evidence_id: str


@app.post('/api/cases/{case_id}/deadlines/{deadline_id}/confirm')
def deadline_confirm(case_id:str,deadline_id:str,data:DeadlineConfirm,user=Depends(staff)):
    def apply(c):
        deadline=domain.item(c,'deadlines',deadline_id)
        domain.evidence(c,[data.evidence_id])
        deadline.update(due_date=data.due_date.isoformat(),evidence_id=data.evidence_id,status='confirmed',confirmed_by=user['name'],reason=data.reason)
        for corr in c['corrections']:
            if corr.get('deadline_id')==deadline_id:
                corr.update(due_date=data.due_date.isoformat(),deadline_evidence_id=data.evidence_id)
                if corr['status']=='approved':
                    corr['status']='reviewed'
                for approval in c['approvals']:
                    if approval.get('correction_id')==corr['id']:
                        approval.update(stale=True,stale_reason='보정 기한 변경')
        for bundle in c['bundles']:
            bundle.update(stale=True,stale_reason='기한 확인·변경')
        for approval in c['approvals']:
            if approval.get('bundle_id'):
                approval.update(stale=True,stale_reason='기한 확인·변경')
    return change(case_id,user,data.expected_version,'deadline.confirmed',apply)


class Submission(Payload):
    bundle_id: str
    receipt_document_id: str
    court_case_number: str


@app.post('/api/cases/{case_id}/submissions')
def submission(case_id:str,data:Submission,user=Depends(staff)):
    from . import filing
    return change(case_id,user,data.expected_version,'submission.recorded',
                  lambda case:filing.record_submission(case,data.model_dump(),user))


class AgentRequest(Payload):
    kind: Literal['case_review','correction']='case_review'


def run_worker(run_id,case,kind):
    with store.db() as con:
        row=con.execute('SELECT body FROM runs WHERE id=?',(run_id,)).fetchone()
        run=json.loads(row['body'])
        run['status']='running'
        con.execute('UPDATE runs SET status=?,body=? WHERE id=?',('running',store.dumps(run),run_id))
    try:
        from .agent_engine import run_analysis
        output=asyncio.run(run_analysis(case,kind=kind))
        if store.get_case(case['id'])['version']!=case['version']:
            output.update(status='stale',error='실행 중 사건 근거가 변경되어 결과가 만료되었습니다.')
        run.update(output)
    except Exception as exc:
        run.update(status='failed',error=f'모델 실행 실패 ({type(exc).__name__}). 기본 업무는 계속할 수 있습니다.')
    run['finished_at']=store.now()
    with store.db() as con:
        con.execute('UPDATE runs SET status=?,body=? WHERE id=?',(run['status'],store.dumps(run),run_id))


@app.post('/api/cases/{case_id}/agent-runs')
def agent_run(case_id:str,data:AgentRequest,user=Depends(staff)):
    case=authorize(case_id,user)
    intake_workflow.require_completed(case)
    if case['version']!=data.expected_version:
        raise store.VersionConflict('현재 사건 버전으로 다시 실행하세요.')
    domain.require(data.kind!='correction' or bool(case['corrections']),'CORRECTION_REQUIRED','보정 요청이 등록된 후 실행합니다.')
    with store.db() as con:
        con.execute('BEGIN IMMEDIATE')
        rows=con.execute('SELECT body FROM runs WHERE case_id=? AND version=?',(case_id,case['version'])).fetchall()
        for row in rows:
            previous=json.loads(row['body'])
            if previous['kind']==data.kind and previous['status'] in ('queued','running','needs_review','manual_review'):
                return {'run_id':previous['id'],'status':previous['status'],'reused':True}
        domain.require(len(rows)<3,'AGENT_BUDGET','현재 근거 버전의 실행 예산 3회를 사용했습니다. 자료·실패 원인을 검토하세요.')
        pending=con.execute("SELECT count(*) FROM runs WHERE status IN ('queued','running')").fetchone()[0]
        domain.require(pending<4,'QUEUE_FULL','작업 대기열이 가득 찼습니다. 잠시 후 실행하세요.')
        run={'id':store.uid('run'),'case_id':case_id,'input_revision':case['input_revision'],'kind':data.kind,'status':'queued','created_at':store.now(),'requested_by':user['name']}
        con.execute('INSERT INTO runs VALUES (?,?,?,?,?)',(run['id'],case_id,case['version'],'queued',store.dumps(run)))
    POOL.submit(run_worker,run['id'],copy.deepcopy(case),data.kind)
    return {'run_id':run['id'],'status':'queued'}


@app.get('/api/agent-runs/{run_id}')
def agent_status(run_id:str,user=Depends(staff)):
    with store.db() as con:
        row=con.execute('SELECT * FROM runs WHERE id=?',(run_id,)).fetchone()
    if not row:
        raise HTTPException(404,'작업이 없습니다.')
    case=authorize(row['case_id'],user)
    return present_run(row,case)


@app.get('/api/system')
async def system(user=Depends(staff)):
    import httpx
    from .model_client import provider_config
    from .corpus import list_sources
    url=os.getenv('OLLAMA_BASE_URL','http://127.0.0.1:11434').rstrip('/')
    status={'available':False,'base_url':url,'model':os.getenv('OLLAMA_MODEL','llama3:latest')}
    try:
        async with httpx.AsyncClient(timeout=3,trust_env=False) as client:
            tags=(await client.get(url+'/api/tags')).json()
            ps=(await client.get(url+'/api/ps')).json()
            status.update(available=True,models=[m['name'] for m in tags.get('models',[])],running=ps.get('models',[]))
    except Exception:
        status['error']='로컬 모델 서버에 연결되지 않습니다.'
    from . import ocr_engine, court_forms, legal_calculator, prompt_registry
    return {'ollama':status,'model':provider_config(),'corpus':list_sources()['stats'],
        'ocr':ocr_engine.status(),'prompts':prompt_registry.catalog(),
        'court_forms':court_forms.catalog()['coverage'],
        'legal_calculator':{'available':True,'version':legal_calculator.CODE_VERSION,
            'policy_version':legal_calculator.policy()['version'],'approval':'lawyer_review_required'},
        'auto_ax':os.getenv('DEBTOFF_AUTO_AX','1')=='1','workflow_rule_count':len(rulebook.RULES),
        'demo_mode':DEMO_MODE,'legal_rules_active':False,
        'limitations':['정식 서식 원본 기반 초안 제공; 관할별 접수요건·서명·제출은 최종 확인',
            '법률계산은 지원 범위의 검토된 입력 기준; 사업소득·별제권 예정부족액 별도 검토',
            'v2 원본 규칙 등록부의 일괄 운영 승인은 비활성',
            '악성코드 검사 공급자, 실메일·문자·법원 제출 자동화 미연결','단일 서버 SQLite 로컬 검증 구성'],
        'metrics':{'max_concurrent_agent_runs':1,'queue_limit':8,'upload_limit_mb':10,'session_hours':8}}


@app.get('/api/legal/search')
async def legal_search(q:str,target:Literal['law','prec']='law',user=Depends(staff)):
    domain.require(1<=len(q)<=200,'QUERY_LENGTH','검색어는 1~200자입니다.')
    from .legal_sources import search_legal
    return await search_legal(q,target=target)


class IntakeInput(BaseModel):
    model_config=ConfigDict(extra='forbid')
    text: str=Field(min_length=20,max_length=50000)


@app.post('/api/intake-runs')
def intake_start(data:IntakeInput,user=Depends(staff)):
    return ax_service.start_intake(data.text,user,POOL)


@app.post('/api/intake-runs/upload')
async def intake_upload(file:UploadFile=File(...),user=Depends(staff)):
    content=await file.read(10_000_001)
    domain.require(0<len(content)<=10_000_000,'FILE_SIZE','회의록 파일은 10MB 이하여야 합니다.')
    filename=Path((file.filename or 'meeting.txt').replace('\\','/')).name[:160]
    pages=await asyncio.to_thread(extract_file,content,Path(filename).suffix.lower())
    text='\n'.join(p['text'] for p in pages)
    domain.require(bool(text.strip()),'UNREADABLE_INTAKE','회의록 문자를 판독하지 못했습니다. 텍스트 또는 문자가 포함된 PDF/DOCX를 입력하세요.')
    return ax_service.start_intake(text,user,POOL,filename)


@app.get('/api/intake-runs/{run_id}')
def intake_result(run_id:str,user=Depends(staff)):
    with store.db() as con:
        row=con.execute('SELECT body FROM intake_runs WHERE id=? AND org_id=?',(run_id,user['org_id'])).fetchone()
    if not row:raise HTTPException(404,'상담 분석 작업을 찾을 수 없습니다.')
    run=json.loads(row['body']);run.pop('input_text',None)
    return run


class AXInput(Payload):
    kind: Literal['case_review','correction']='case_review'


@app.post('/api/cases/{case_id}/ax-runs')
def ax_start(case_id:str,data:AXInput,user=Depends(staff)):
    case=authorize(case_id,user)
    if case['version']!=data.expected_version:raise store.VersionConflict('사건이 변경되었습니다. 새로고침 후 분석하세요.')
    return ax_service.schedule(case,user,POOL,data.kind)


@app.get('/api/cases/{case_id}/ax-runs')
def ax_list(case_id:str,user=Depends(staff)):
    return {'runs':ax_service.case_runs(authorize(case_id,user),20)}


@app.get('/api/ax-runs/{run_id}')
def ax_result(run_id:str,user=Depends(staff)):
    with store.db() as con:row=con.execute('SELECT * FROM ax_runs WHERE id=?',(run_id,)).fetchone()
    if not row:raise HTTPException(404,'분석 작업을 찾을 수 없습니다.')
    return ax_service.present(row,authorize(row['case_id'],user))


class FindingReview(Reason):
    period: str | None=Field(default=None,max_length=500)


@app.post('/api/cases/{case_id}/ax-runs/{run_id}/findings/{finding_id}/{decision}')
def finding_review(case_id:str,run_id:str,finding_id:str,decision:Literal['apply','dismiss'],data:FindingReview,user=Depends(staff)):
    authorize(case_id,user)
    saved=ax_service.apply_finding(case_id,run_id,finding_id,user,data.expected_version,data.reason,data.period,decision=='dismiss')
    if decision=='apply':
        ax_service.maybe_schedule(saved,user,POOL,trigger='분석 제안 반영')
    return visible(saved,user)


class CandidateReview(Reason):
    decision: Literal['accept','correct','reject']
    value: StrictStr | StrictInt | StrictFloat | StrictBool | list | dict | None=None


@app.post('/api/cases/{case_id}/extraction-candidates/{candidate_id}/review')
def candidate_review(case_id:str,candidate_id:str,data:CandidateReview,user=Depends(staff)):
    def apply(c):
        candidate=domain.item(c,'extraction_candidates',candidate_id)
        source_review.validate_candidate(candidate,data.decision,data.value)
        domain.invalidate(c,'추출값 검토 반영')
        document=next((doc for doc in c.get('documents',[]) if doc['id']==candidate.get('document_id')),None)
        source_review.apply_candidate(candidate,data.decision,data.value,data.reason,user,document=document,case=c)
        # Preserve the authored document versions and their review metadata.
        # change() queues the evidence/analysis/document pipeline for this edit;
        # a synchronous generic draft would replace that verified writing with
        # unreviewed placeholder prose before the pipeline has run.
    return change(case_id,user,data.expected_version,'candidate.reviewed',apply)


class DraftReview(Reason):
    decision: Literal['approve','request_changes']


@app.post('/api/cases/{case_id}/drafts/{draft_id}/review')
def draft_review(case_id:str,draft_id:str,data:DraftReview,user=Depends(staff)):
    domain.require(user['role']=='lawyer','LAWYER_ONLY','1차 서류 검토는 변호사 계정에서 완료하세요.')
    def apply(c):
        draft=domain.item(c,'drafts',draft_id)
        domain.require(not draft['stale'] and draft['input_revision']==c['input_revision'],'STALE_DRAFT','자료 변경으로 만료된 초안입니다. 최신 초안을 검토하세요.')
        domain.require(draft['content_hash']==store.digest({k:draft[k] for k in ('sections','source_refs','input_revision')}),'DRAFT_HASH','초안 내용과 저장 해시가 다릅니다.')
        draft.update(status='reviewed' if data.decision=='approve' else 'changes_requested',review={'decision':data.decision,'reason':data.reason,'actor':user['name'],'at':store.now(),'scope':'1차 초안 검토; 법원 제출 승인과 별도'})
        c['automation'].update(stage='reviewed' if data.decision=='approve' else 'changes_requested',label='변호사 1차 검토 완료' if data.decision=='approve' else '변호사 수정 요청')
    return change(case_id,user,data.expected_version,'draft.reviewed',apply)


class IntakeReview(Reason):
    client_name: str=Field(min_length=1,max_length=60)
    region: str=Field(default='',max_length=200)
    court_id: str
    case_type: Literal['personal_rehabilitation','bankruptcy_review','other','unknown']


@app.post('/api/cases/{case_id}/intake-review')
def intake_review(case_id:str,data:IntakeReview,user=Depends(staff)):
    court=next((c for c in registry()['courts'] if c['id']==data.court_id),None)
    domain.require(court is not None or data.court_id=='unknown','COURT_UNKNOWN','관할 후보를 등록부에서 선택하세요.')
    def apply(c):
        domain.invalidate(c,'담당자 사건 분류 검토')
        identity_changed=c['client_name']!=data.client_name
        previous_name=c['client_name']
        if identity_changed:
            for doc in c['documents']:
                doc.update(status='quarantined',person_confirmed=False,public_status='명의 변경에 따른 재확인 필요',automated_check={'document_id':doc['id'],'coverage_status':'identity_conflict','reason':'내담자 성명이 변경되어 자료 귀속을 다시 확인해야 합니다.','missing':['새 성명과 자료 대상자 대조']})
            for request in c['requests']:
                if request.get('document_ids'):
                    request.update(status='needs_more',public_review_note='자료 대상자를 다시 확인하고 있습니다.')
            for field in ('extraction_candidates','facts','messages'):
                for record in c.get(field,[]):
                    record.update(status='quarantined',quarantine_reason='내담자 성명 변경 후 자료 귀속 재확인')
            if c.get('consultation'):c['consultation']['status']='quarantined'
            c['summary']=data.client_name+' · 성명 변경으로 기존 자료 귀속 재확인 필요'
        c.update(client_name=data.client_name,region=data.region or None,court_id=data.court_id,court_name=court['name'] if court else '관할 후보 확인 필요',case_type=data.case_type)
        c['case_type_label']={'personal_rehabilitation':'개인회생','bankruptcy_review':'파산 검토','other':'다른 사건유형','unknown':'사건유형 확인 필요'}[data.case_type]
        c['title']=c['client_name']+' · '+c['case_type_label']
        c['intake_review']={'actor':user['name'],'reason':data.reason,'at':store.now()}
        if identity_changed:
            c.setdefault('extraction_candidates',[]).append({'id':store.uid('candidate'),'key':'client_name','label':'내담자 성명','original_value':previous_name,'value':data.client_name,'status':'accepted','origin':'human_correction','source_type':'human_review','source_id':'intake_review','quote':data.reason,'review':c['intake_review']})
        for candidate in c.get('extraction_candidates',[]):
            if candidate['key']=='client_name' and candidate.get('status') not in ('rejected','quarantined'):
                candidate.update(original_value=candidate.get('original_value',candidate['value']),value=data.client_name,status='accepted',origin='human_correction',review=c['intake_review'])
        if c['case_type'] in ('personal_rehabilitation','bankruptcy_review'):
            ax_service.ensure_requests(c,rulebook.evaluate_case(c)['required_documents'])
        # Classification edits invalidate old outputs but do not author replacements.
        # The scheduled pipeline checks extraction and review completion first.
    return change(case_id,user,data.expected_version,'intake.reviewed',apply)


@app.get('/api/cases/{case_id}/drafts/{draft_id}/download')
def draft_download(case_id:str,draft_id:str,format:Literal['html','docx']='docx',user=Depends(staff)):
    case=authorize(case_id,user);draft=domain.item(case,'drafts',draft_id)
    store.access(user,case_id,'draft.download:'+draft_id)
    disposition = download_names.disposition(download_names.filename(case, draft['title'], format))
    if format=='html':
        return HTMLResponse(drafting.to_html(case,draft),headers={'Content-Disposition':disposition,'Content-Security-Policy':"default-src 'none'; style-src 'unsafe-inline'"})
    return Response(drafting.to_docx(case,draft),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':disposition})


@app.get('/api/rulebook')
def workflow_rules(user=Depends(staff)):
    return rulebook.definition()


@app.post('/api/cases/{case_id}/automation/run')
def automation_run(case_id: str, data: Payload, user=Depends(staff)):
    case = authorize(case_id, user)
    if case['version'] != data.expected_version:
        raise store.VersionConflict('사건이 변경되었습니다. 새로고침 후 다시 실행하세요.')
    return ax_service.schedule(case, user, POOL, trigger='자동 처리 재시도')


@app.post('/api/cases/{case_id}/notifications/{notice_id}/read')
def notification_read(case_id: str, notice_id: str, data: Payload, user=Depends(current_user)):
    def apply(case):
        notice = domain.item(case, 'notifications', notice_id)
        domain.require(notice.get('audience') == ('client' if user['role'] == 'client' else 'staff'),
                       'NOTICE_ACCESS', '이 알림에 접근할 수 없습니다.')
        notifications.mark_read(notice, user)
    return change(case_id, user, data.expected_version, 'notification.read', apply)


class CourtOutcome(Payload):
    outcome: Literal['approved', 'correction', 'rejected']
    reason: str = Field(min_length=5, max_length=4000)
    bundle_id: str = Field(min_length=1, max_length=100)
    source_document_id: str = Field(min_length=1, max_length=100)
    decision_type: Literal['initial_plan_approval', 'initial_plan_denial', 'application_dismissal',
                           'commencement', 'correction_order', 'discharge', 'other'] | None = None
    decision_date: str | None = Field(default=None, pattern=r'^\d{4}-\d{2}-\d{2}$')
    decision_quote: str | None = Field(default=None, min_length=5, max_length=1500)


class DocumentMetadata(Reason):
    metadata: dict


@app.post('/api/cases/{case_id}/documents/{doc_id}/metadata')
def document_metadata(case_id: str, doc_id: str, data: DocumentMetadata, user=Depends(staff)):
    allowed = {'institution','account_key','period_start','period_end','issued_at','certificate_type',
               'address_history','tax_scope','jurisdiction_scope','person_number_display','covers_all_accounts'}
    domain.require(set(data.metadata) <= allowed and all(type(v) in (str, bool) and len(str(v)) <= 300 for v in data.metadata.values()),
                   'METADATA_SCOPE', '서류 기관·계좌·기간·발급옵션 형식을 확인하세요.')
    for key in ('period_start', 'period_end', 'issued_at'):
        if data.metadata.get(key):
            try: date.fromisoformat(data.metadata[key])
            except (ValueError, TypeError): raise HTTPException(422, '날짜는 YYYY-MM-DD 형식입니다.')
    def apply(case):
        doc = domain.item(case, 'documents', doc_id)
        domain.invalidate(case, '서류 범위 판독값 수정')
        doc.update(document_metadata=data.metadata, metadata_review={'actor':user['id'],'reason':data.reason,'at':store.now()})
        if doc.get('request_id'):
            request = domain.item(case, 'requests', doc['request_id'])
            # Correcting an archived original does not restore an excluded request.
            if request.get('status') not in automation.INACTIVE and not request.get('no_longer_required'):
                request['status'] = 'received'
    return change(case_id, user, data.expected_version, 'document.metadata', apply)


class DocumentScopes(Payload):
    application_date: date | None = None
    financial_accounts: list[dict] = Field(default_factory=list, max_length=100)
    document_scopes: list[dict] = Field(default_factory=list, max_length=300)
    creditors: list[dict] = Field(default_factory=list, max_length=500)
    insurance_policies: list[dict] = Field(default_factory=list, max_length=100)
    employers: list[dict] = Field(default_factory=list, max_length=100)


@app.post('/api/cases/{case_id}/document-scopes')
def document_scopes(case_id: str, data: DocumentScopes, user=Depends(staff)):
    values = data.model_dump(mode='json', exclude={'expected_version'}, exclude_unset=True)
    domain.require(len(store.dumps(values)) <= 100000, 'SCOPE_SIZE', '요청 범위가 너무 큽니다.')
    allowed = {'id','catalog_id','institution','account_number','account_key','account_label','purpose',
               'salary_account','status','name','issuer','period_start','period_end','period_label','issuance_options'}
    for name, rows in values.items():
        if name == 'application_date':
            continue
        for row in rows:
            domain.require(set(row) <= allowed, 'SCOPE_FIELDS', '허용된 기관·계좌·기간 항목만 수정할 수 있습니다.')
            for key, value in row.items():
                if key == 'issuance_options' and isinstance(value, dict):
                    option_keys = {'freshness_months','certificate_type','address_history','tax_scope',
                                   'jurisdiction_scope','person_number_display','include_closed_accounts'}
                    domain.require(set(value) <= option_keys, 'ISSUANCE_OPTIONS', '지원하는 발급 옵션만 지정하세요.')
                    if 'freshness_months' in value:
                        domain.require(type(value['freshness_months']) is int and 1 <= value['freshness_months'] <= 2,
                                       'ISSUANCE_FRESHNESS', '일반 요청의 발급 유효기간은 1~2개월입니다. 예외는 개별 법원 요구로 기록하세요.')
                domain.require((key == 'issuance_options' and isinstance(value, dict) and
                    all(type(v) in (str, int, bool) and len(str(v)) <= 300 for v in value.values())) or
                    (key != 'issuance_options' and type(value) in (str, bool) and len(str(value)) <= 300),
                    'SCOPE_VALUES', '기관·계좌·기간 값의 형식을 확인하세요.')
            if row.get('period_start') or row.get('period_end'):
                try:
                    domain.require(date.fromisoformat(row['period_start']) <= date.fromisoformat(row['period_end']), 'SCOPE_DATES', '시작일·종료일을 확인하세요.')
                except (KeyError, ValueError): raise HTTPException(422, '기간 시작일·종료일은 YYYY-MM-DD 형식입니다.')
    def apply(case):
        domain.invalidate(case, '기관·계좌·기간 범위 수정')
        case.update(values)
        ax_service.ensure_requests(case, rulebook.evaluate_case(case)['required_documents'])
    return change(case_id, user, data.expected_version, 'document.scopes', apply)


@app.post('/api/cases/{case_id}/requests/{request_id}/withdraw')
def withdraw_request(case_id: str, request_id: str, data: Reason, user=Depends(staff)):
    def apply(case):
        req = domain.item(case, 'requests', request_id)
        domain.require(req['status'] not in automation.INACTIVE and not req.get('no_longer_required'),
                       'REQUEST_WITHDRAWN', '이미 제외된 요청입니다.')
        reason = data.reason.strip()
        domain.require(len(reason) >= 5, 'REASON_REQUIRED', '제외하는 이유를 5자 이상 기록하세요.')
        domain.invalidate(case, '불필요 서류 요청 철회')
        timestamp, previous = store.now(), req['status']
        req.update(status='withdrawn', no_longer_required=True, manual_override=True,
                   withdrawal_reason=reason, withdrawn_at=timestamp, withdrawn_by=user['id'],
                   version=req.get('version', 1) + 1)
        case.setdefault('request_history', []).append({'id':store.uid('reqevt'),'request_id':request_id,
            'action':'withdrawn','reason':reason,'actor':user['id'],'at':timestamp,
            'previous_status':previous,'scope_key':req.get('scope_key')})
        automation.sync_request_notifications(case)
    return change(case_id, user, data.expected_version, 'request.withdrawn', apply)


@app.post('/api/cases/{case_id}/court-outcomes')
def court_outcome(case_id: str, data: CourtOutcome, user=Depends(staff)):
    return change(case_id, user, data.expected_version, 'court_outcome.recorded',
                  lambda case: automation.record_outcome(case, data.model_dump(exclude={'expected_version'}), user))


@app.get('/api/cases/{case_id}/automation/evidence')
def automation_evidence(case_id: str, user=Depends(staff)):
    case = authorize(case_id, user)
    return {key: case.get(key, []) for key in ('structured_data', 'strategy_analyses', 'verification_runs', 'court_outcomes', 'request_history')}


@app.get('/api/court-request-rules')
def court_request_rules(user=Depends(staff)):
    path = ROOT / 'data/court_request_rules.json'
    return json.loads(path.read_text(encoding='utf-8'))


def _legal_watch_view():
    from . import legal_watch, reasoning_cache
    state = legal_watch.status()
    if LEGAL_WATCH_JOB and not LEGAL_WATCH_JOB.done():
        state['status'] = 'running'
    elif LEGAL_WATCH_JOB and LEGAL_WATCH_JOB.exception():
        state.update(status='failed', message='공식 근거 확인을 완료하지 못했습니다. 마지막 정상 근거와 확인 이력을 유지합니다.')
    metrics = reasoning_cache.operational_metrics()
    return {**state, 'scheduler_enabled': os.getenv('DEBTOFF_LEGAL_WATCH', '0') == '1',
        'budget': {**state.get('budget', {}),
        'daily_limit': metrics['daily_limit'], 'used': metrics['used'], 'cached': metrics['cached']}}


@app.get('/api/legal-watch')
def legal_watch_status(user=Depends(staff)):
    return _legal_watch_view()


def _manual_legal_watch():
    from . import legal_watch
    asyncio.run(legal_watch.run_once(force=True))
    _refresh_watch_cases()


@app.post('/api/legal-watch/run')
def legal_watch_run(user=Depends(staff)):
    global LEGAL_WATCH_JOB
    with LEGAL_WATCH_LOCK:
        if LEGAL_WATCH_JOB and not LEGAL_WATCH_JOB.done():
            return {'status': 'running', 'reused': True}
        LEGAL_WATCH_JOB = COLLECTION_POOL.submit(_manual_legal_watch)
    store.access(user, None, 'legal_watch.run')
    return {'status': 'queued'}


class LegalWatchPolicy(BaseModel):
    model_config = ConfigDict(extra='forbid')
    enabled: StrictBool
    interval_hours: StrictInt = Field(ge=24, le=720)
    max_sources: StrictInt = Field(ge=1, le=20)
    source_ids: list[StrictStr] = Field(max_length=100)
    daily_reasoning_limit: StrictInt = Field(ge=1, le=100)


@app.post('/api/legal-watch/policy')
def legal_watch_policy(data: LegalWatchPolicy, user=Depends(staff)):
    domain.require(user['role'] == 'lawyer', 'LAWYER_REQUIRED', '변호사만 확인 범위와 비용 한도를 변경할 수 있습니다.')
    from . import legal_watch, reasoning_cache
    try:
        legal_watch.configure(data.model_dump(exclude={'daily_reasoning_limit'}))
        reasoning_cache.configure({'daily_call_limit': data.daily_reasoning_limit})
    except ValueError:
        raise HTTPException(422, '등록된 공식 출처와 확인 범위를 선택하세요.')
    store.access(user, None, 'legal_watch.policy:' + store.dumps(data.model_dump()))
    return _legal_watch_view()


@app.get('/api/prompts')
def model_prompts(user=Depends(staff)):
    from .prompt_registry import catalog
    return catalog(include_text=True)


@app.get('/api/knowledge/sources')
def knowledge_sources(user=Depends(staff)):
    from .corpus import list_sources
    return list_sources()


@app.get('/api/knowledge/search')
def knowledge_search(q:str,court_id:str|None=None,user=Depends(staff)):
    domain.require(1<=len(q)<=200,'QUERY_LENGTH','검색어는 1~200자입니다.')
    from .corpus import search
    return {'results':search(q,court_id=court_id,limit=12)}


@app.get('/api/knowledge/sources/{source_id}')
def knowledge_source(source_id:str,user=Depends(staff)):
    from .corpus import source_detail
    source=source_detail(source_id)
    if not source:raise HTTPException(404,'공식 출처를 찾을 수 없습니다.')
    return source


@app.post('/api/knowledge/sources/{source_id}/refresh')
def refresh_knowledge_source(source_id: str, user=Depends(staff)):
    from .corpus import seeds
    if source_id not in {source['id'] for source in seeds()}:
        raise HTTPException(404, '등록된 공식 출처를 찾을 수 없습니다.')
    return _start_collection(user, [source_id])


def collect_worker(run):
    from .corpus import ingest_all
    from .legal_watch import assess_collected_sources
    run['status']='running';ax_service._save('collection_runs',run)
    try:
        result=asyncio.run(ingest_all(source_ids=run.get('source_ids')))
        assessed = asyncio.run(assess_collected_sources(source_ids=run.get('source_ids')))
        run.update(status='completed',result=result)
        if assessed.get('status') == 'already_running':
            run['pending_legal_assessment'] = True
        else:
            _refresh_watch_cases()
    except Exception as exc:
        run.update(status='failed',error=f'공식 출처 수집 실패 ({type(exc).__name__})')
    run['finished_at']=store.now();ax_service._save('collection_runs',run)


@app.post('/api/knowledge/ingest')
def collect_start(user=Depends(staff)):
    return _start_collection(user)


def _start_collection(user, source_ids=None):
    with store.db() as con:
        con.execute('BEGIN IMMEDIATE')
        old=con.execute("SELECT id,status,body FROM collection_runs WHERE status IN ('queued','running') LIMIT 1").fetchone()
        if old:
            active_sources = json.loads(old['body']).get('source_ids')
            if active_sources is None or (source_ids is not None and set(source_ids) <= set(active_sources)):
                return {'run_id':old['id'],'status':old['status']}
            raise HTTPException(409, '다른 공식 자료를 수집 중입니다. 완료 후 이 원문을 다시 수집하세요.')
        run={'id':store.uid('collection'),'org_id':user['org_id'],'status':'queued','created_at':store.now(), 'source_ids': source_ids}
        con.execute('INSERT INTO collection_runs VALUES (?,?,?,?)',(run['id'],user['org_id'],run['status'],store.dumps(run)))
    COLLECTION_POOL.submit(collect_worker,run)
    return {'run_id':run['id'],'status':'queued'}


@app.get('/api/knowledge/runs/{run_id}')
def collect_result(run_id:str,user=Depends(staff)):
    with store.db() as con:row=con.execute('SELECT body FROM collection_runs WHERE id=? AND org_id=?',(run_id,user['org_id'])).fetchone()
    if not row:raise HTTPException(404,'수집 작업을 찾을 수 없습니다.')
    return json.loads(row['body'])


from .extended_routes import attach as attach_document_workflow
attach_document_workflow(app, staff, authorize, change)
from .document_review_routes import attach as attach_document_review
attach_document_review(app, staff, authorize, change)
source_review.attach(app, staff, change)
from .filing_routes import attach as attach_filing_workflow
attach_filing_workflow(app, staff, authorize, change, extract_file)

from .portal_routes import register as attach_customer_portal


attach_customer_portal(app, current_user, staff, authorize, change, visible)

from .archive_import import attach as attach_archive_import
attach_archive_import(app, current_user, authorize, extract_file, _commit_prepared_documents)
