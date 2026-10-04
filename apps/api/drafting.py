"""Automatically compose first drafts from attributed candidates, leaving unknowns explicit."""
import html
import io
import json
from . import store, rulebook


def value_text(value):
    if value is None or value=='':
        return '미확인'
    if isinstance(value, list):
        return '\n'.join('· ' + value_text(item) for item in value if item is not None)
    if isinstance(value, dict):
        names = {'title': '서류명', 'period': '요청 기간', 'status': '상태', 'due_date': '제출 기한',
                 'institution': '기관', 'issuer': '발급처', 'account_masked': '계좌', 'reason': '사유',
                 'message': '안내', 'value': '확인 내용', 'label': '항목', 'source_ids': '근거',
                 'period_start': '시작일', 'period_end': '종료일', 'issuance_options': '발급 옵션',
                 'evidence_ids': '증빙', 'name': '명칭', 'amount': '금액', 'quote': '원문 근거'}
        statuses = {'requested': '제출 요청', 'received': '제출 완료', 'fulfilled': '확인 완료',
                    'needs_more': '보완 필요', 'withdrawn': '요청 철회', 'verified': '검증 완료',
                    'pending': '확인 대기', 'unknown': '미확인', 'confirmed': '확인됨'}
        return '\n'.join(f"{names.get(key, key.replace('_', ' '))}: " +
            (statuses.get(item, item) if key == 'status' and isinstance(item, str) else value_text(item))
            for key, item in value.items() if item is not None and item != '')
    if type(value) is bool:
        return '해당' if value else '해당 없음'
    if isinstance(value, str) and value.strip().startswith(('{', '[')):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, (dict, list)):
                return value_text(parsed)
        except (ValueError, TypeError):
            pass
    return str(value)


def narrative_citations(draft, section):
    sources = {source['id']: source for source in draft.get('narrative_sources', [])}
    citations, seen = [], set()
    for paragraph in section.get('grounded_paragraphs', []):
        for quote in paragraph.get('quotes', []):
            source = sources.get(quote.get('source_id'), {})
            key = (quote.get('source_id'), quote.get('quote'))
            if source.get('kind') == 'public_legal_source' and key not in seen:
                seen.add(key)
                citations.append({'title': source.get('title') or source.get('public_source_id'),
                    'url': source.get('url'), 'quote': quote['quote'], 'effective_date': source.get('effective_date')})
    return citations


def generate(case,trigger='자료 분석'):
    evaluation=rulebook.evaluate_case(case)
    case['rule_evaluation']=evaluation
    excluded={d['id'] for d in case['documents'] if d.get('status') in ('rejected','quarantined','superseded')}
    candidates=[c for c in case.get('extraction_candidates',[]) if c.get('status') not in ('rejected','quarantined','superseded') and c.get('document_id') not in excluded]
    chosen={}
    def rank(candidate):
        return (candidate.get('status')=='accepted',candidate.get('source_type') not in ('meeting','party_statement'))
    for candidate in candidates:
        if candidate.get('status')=='rejected' or candidate.get('value') is None:
            continue
        key=candidate['key']
        if key not in chosen or rank(candidate)>=rank(chosen[key]):
            chosen[key]=candidate
    for fact in case['facts']:
        if fact.get('status')=='confirmed' and fact.get('value') is not None and fact.get('evidence_ids') and all(any(d['id']==eid and d.get('status')=='verified' for d in case['documents']) for eid in fact['evidence_ids']):
            chosen[fact['key']]={'key':fact['key'],'label':fact['label'],'value':fact['value'],'status':'confirmed','source_ids':fact['evidence_ids'],'quote':fact.get('reason','')}
    conflicts=[]
    for key in chosen:
        values={json.dumps(x.get('value'),ensure_ascii=False,sort_keys=True) for x in candidates if x['key']==key and x.get('status')!='rejected' and x.get('value') is not None}
        if len(values)>1:
            conflicts.append(key)
    sections=[]
    rehabilitation=case.get('case_type','personal_rehabilitation')=='personal_rehabilitation'
    targets=('application','creditors','assets','income_expense','statement','repayment_plan','attachments') if rehabilitation else ('application','creditors','assets','income_expense','statement','attachments')
    for target in targets:
        fields=[]
        for f in evaluation['factor_definitions']:
            if f['draft_target']!=target: continue
            c=chosen.get(f['key'],{})
            fields.append({'key':f['key'],'label':f['label'],'value':c.get('value'),'source_ids':c.get('source_ids',[c.get('document_id') or c.get('source_id')]) if c else [],'status':'conflict' if f['key'] in conflicts else c.get('status','unknown')})
        if target=='application':
            content=f"{case['client_name']}님의 {case.get('region') or '지역 미확인'} 상담을 바탕으로 작성한 {case.get('case_type_label','개인회생 검토')} 초안입니다. 관할 후보: {case['court_name']}. 관할과 신청 방향은 변호사가 검토합니다."
        elif target=='creditors':
            content='채권자별 기관·원금·이자·기준일·양도/대위 여부를 증빙과 대조합니다. 총액만 확보된 경우 채권자별 금액을 임의 배분하지 않습니다.'
        elif target=='assets':
            content='제출자료에서 관찰된 재산 후보를 정리합니다. 자료가 없다는 이유로 무재산으로 기재하지 않습니다.'
        elif target=='income_expense':
            content='상담 진술과 증빙 추출값을 출처별로 구분했습니다. 실제 지출은 법률상 인정 생계비와 구분하여 검토합니다.'
        elif target=='statement':
            content='채무 발생 경위, 상환이 어려워진 과정, 현재 소득과 변제 가능성을 근거에 따라 작성하는 단계입니다. 작성·검증이 완료된 진술서가 여기에 반영됩니다.'
            if case.get('consultation',{}).get('status')=='quarantined':content='내담자 성명 변경으로 기존 상담 진술의 귀속을 재확인해야 합니다. 현재 초안에는 해당 진술을 반영하지 않았습니다.'
        elif target=='repayment_plan':
            content='소득·지출·재산·채권 후보를 바탕으로 변제계획 검토 입력을 정리했습니다. 법률상 생계비·기간·청산가치·우선권·배분·비용을 검토하기 전 확정 변제금은 산출하지 않습니다.'
            fields=[{'key':k,'label':label,'value':chosen.get(k,{}).get('value'),'status':chosen.get(k,{}).get('status','unknown'),'source_ids':[chosen.get(k,{}).get('document_id')]} for k,label in [('monthly_income','월 소득 후보'),('living_expenses','실제 월 지출 후보'),('total_debt','총 채무 후보'),('assets_total','재산 총액 후보')]]
        else:
            content='\n'.join(f"{r['title']} / 범위: {r['period']} / 상태: {r['status']}" for r in case['requests'] if r.get('status') not in ('withdrawn','superseded','cancelled')) or '자료 요청안 생성 전'
        title=evaluation['draft_targets'][target]
        if not rehabilitation and target=='application':title='상담 및 절차 방향 검토 초안'
        sections.append({'id':target,'title':title,'content':content,'fields':fields})
    missing=[f['label'] for f in evaluation['factor_definitions'] if f['key'] in evaluation['missing_factors']]
    missing.extend(f"{next((f['label'] for f in evaluation['factor_definitions'] if f['key']==key),key)}: 출처 간 값이 달라 검토 필요" for key in conflicts)
    for old in case.setdefault('drafts',[]):
        old.update(stale=True,stale_reason=trigger+' 후 초안 갱신')
    draft={'id':store.uid('draft'),'title':'1차 신청서류 검토 패키지','status':'draft','version':len(case['drafts'])+1,'input_revision':case['input_revision'],'created_at':store.now(),'trigger':trigger,'sections':sections,'missing_fields':missing,'conflicting_factors':conflicts,'source_refs':[{'document_id':c.get('document_id'),'source_id':c.get('source_id'),'quote':c.get('quote',''),'page':c.get('page')} for c in candidates if c.get('status')!='rejected'],'stale':False,'scope':'변호사 검토용 1차 초안 · 법원 제출 승인 전','review':None}
    if not rehabilitation:
        draft['title']='상담 및 절차 검토 패키지'
    draft['content_hash']=store.digest({'sections':sections,'source_refs':draft['source_refs'],'input_revision':draft['input_revision']})
    from .automation import _features
    draft['generation_features'] = _features(case, generated=draft)
    case['drafts'].append(draft)
    case['automation']={'stage':'lawyer_review','label':'1차 서류가 준비되었습니다. 변호사 검토 대기','draft_id':draft['id'],'generated_at':store.now(),'missing_count':len(missing),'conflict_count':len(conflicts),'source_count':len(case['documents'])}
    return draft


def to_html(case,draft):
    e=lambda v:html.escape(value_text(v))
    sections=[]
    for s in draft['sections']:
        rows=''.join(f"<tr><td>{e(f['label'])}</td><td>{e(f['value'])}</td><td>{e(f['status'])}</td><td>{e(', '.join(str(v) for v in f.get('source_ids',[]) if v))}</td></tr>" for f in s['fields'])
        citations = ''.join(f"<blockquote>{e(ref['title'])} · {e(ref['effective_date'] or '시행일 별도 확인')}<br>{e(ref['quote'])}<br>{e(ref['url'])}</blockquote>" for ref in narrative_citations(draft, s))
        sections.append(f"<section><h2>{e(s['title'])}</h2><p>{e(s['content']).replace(chr(10),'<br>')}</p>{citations}<table><tr><th>항목</th><th>추출값/초안</th><th>확인상태</th><th>출처</th></tr>{rows}</table></section>")
    missing=''.join('<li>'+e(x)+'</li>' for x in draft['missing_fields'])
    return f'<!doctype html><html lang="ko"><meta charset="utf-8"><title>빚오프 1차 서류</title><style>body{{font:14px sans-serif;line-height:1.8;max-width:980px;margin:40px auto}}section{{page-break-before:always}}table{{width:100%;border-collapse:collapse}}td,th{{padding:8px;border:1px solid #ddd}}.notice{{background:#fff4dd;padding:18px}}</style><h1>{e(case["client_name"])} · 1차 신청서류</h1><p class="notice">{e(draft["scope"])} · 버전 {draft["version"]} · {"만료된 초안" if draft["stale"] else e(draft["status"])}</p><h2>변호사 확인 항목</h2><ul>{missing}</ul>'+''.join(sections)+f'<p>출력 해시 {draft["content_hash"]}</p></html>'


def to_docx(case,draft):
    from docx import Document
    from docx.shared import Pt
    doc=Document()
    doc.styles['Normal'].font.name='맑은 고딕'
    doc.styles['Normal'].font.size=Pt(10)
    doc.add_heading(case['client_name']+' · '+draft['title'],0)
    doc.add_paragraph(draft['scope']+' / 버전 '+str(draft['version'])+' / '+('만료된 초안' if draft['stale'] else draft['status']))
    doc.add_heading('변호사 확인 항목',1)
    for missing in draft['missing_fields']:doc.add_paragraph(missing,style='List Bullet')
    for s in draft['sections']:
        doc.add_page_break();doc.add_heading(s['title'],1);doc.add_paragraph(s['content'])
        for ref in narrative_citations(draft, s):
            doc.add_paragraph(str(ref['title']) + ' · ' + (ref['effective_date'] or '시행일 별도 확인') + '\n' + ref['quote'] + '\n' + str(ref['url']))
        table=doc.add_table(rows=1,cols=4);table.style='Table Grid'
        for cell,value in zip(table.rows[0].cells,['항목','추출값','확인상태','원문근거']):cell.text=value
        for f in s['fields']:
            for cell,value in zip(table.add_row().cells,[f['label'],value_text(f['value']),f['status'],', '.join(str(x) for x in f.get('source_ids',[]) if x)]):cell.text=value
    doc.add_page_break();doc.add_heading('추출값 원문 근거',1)
    for ref in draft['source_refs']:
        doc.add_paragraph(str(ref.get('document_id') or ref.get('source_id') or '')+' / '+str(ref.get('page') or '쪽 미확인')+'\n'+ref.get('quote',''))
    doc.add_paragraph('출력 해시: '+draft['content_hash'])
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()
