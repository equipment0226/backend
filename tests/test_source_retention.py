"""Late evidence in large accepted files must not disappear before extraction."""
import io
import unittest
from unittest.mock import patch
from docx import Document
from apps.api import main


class SourceRetentionTests(unittest.TestCase):
    def test_text_keeps_late_account_evidence(self):
        text = '거래 내역\n' * 10001 + '기말잔액: 7,350,000원'
        pages = main.extract_file(text.encode('utf-8'), '.txt')
        self.assertEqual(pages[0]['text'], text)
        self.assertTrue(pages[0]['text'].endswith('기말잔액: 7,350,000원'))

    def test_docx_keeps_late_creditor_table(self):
        doc = Document()
        doc.add_paragraph('참고자료 ' * 12001)
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = '채권자: 두번째테스트은행'
        row.cells[1].text = '원금: 4,500,000원'
        stream = io.BytesIO()
        doc.save(stream)
        text = main.extract_file(stream.getvalue(), '.docx')[0]['text']
        self.assertGreater(len(text), 60000)
        self.assertIn('채권자: 두번째테스트은행 | 원금: 4,500,000원', text)

    def test_pdf_keeps_every_returned_page(self):
        import fitz
        pdf = fitz.open()
        pdf.new_page()
        raw = pdf.tobytes()
        pdf.close()
        pages = [{'page': 1, 'text': '첫 거래\n' * 13000}, {'page': 2, 'text': '별도 계좌 잔액: 950,000원'}]
        with patch('apps.api.ocr_engine.extract_pages', return_value=pages):
            self.assertEqual(main.extract_file(raw, '.pdf'), pages)
