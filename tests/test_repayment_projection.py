"""Default term and source-bound, explicitly unapproved repayment scenarios."""
import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from apps.api import legal_calculator as calc


class RepaymentProjectionTests(unittest.TestCase):
    def setUp(self):
        self.inputs = calc.example_payload()
        rows = [self.inputs['income']] + self.inputs['assets'] + self.inputs['creditors']
        self.case = {'id': 'synthetic', 'court_id': 'CT01', 'input_revision': 1,
            'case_type': 'personal_rehabilitation', 'issues': [],
            'documents': [{'id': eid, 'status': 'verified'} for row in rows for eid in row['evidence_ids']]}
        unknown = copy.deepcopy(self.inputs)
        for key in ('recognized_household_size', 'additional_living_cost', 'months', 'prepaid_months',
                    'monthly_trustee_fee', 'preapproval_costs_paid', 'objection', 'living_cost_mode'):
            unknown[key] = None
        unknown['decisions'] = {}
        for asset in unknown['assets']:
            for key in ('secured_deduction', 'exempt_deduction', 'disposal_cost'):
                asset[key] = None
        self.packet = {'inputs': unknown, 'form_values': {'household_size': 1},
            'origins': {'household_size': {'source_ids': ['fixture-household']}},
            'errors': [], 'source_signature': 'synthetic-originals-v1'}
        self.mapper = patch('apps.api.evidence_mapping.build', side_effect=lambda _: copy.deepcopy(self.packet)).start()
        self.addCleanup(patch.stopall)

    def test_missing_and_null_months_default_without_mutating_input(self):
        for mode in ('absent', None):
            inputs = copy.deepcopy(self.inputs)
            if mode == 'absent':
                inputs.pop('months')
            else:
                inputs['months'] = None
            before = copy.deepcopy(inputs)
            result = calc.calculate_legal(self.case, inputs)
            self.assertEqual(result['summary']['months'], 36)
            self.assertEqual(result['defaulted_inputs'][0]['source_id'], 'statute-611')
            self.assertEqual(inputs, before)

    def test_explicit_period_including_invalid_zero_is_not_replaced(self):
        for months in (24, 60, 0):
            inputs = {**self.inputs, 'months': months}
            result = calc.calculate_legal(self.case, inputs)
            self.assertEqual(result['inputs']['months'], months)
            self.assertEqual(result['defaulted_inputs'], [])
            if months == 0:
                self.assertEqual(result['schedule'], [])

    def test_missing_judgments_produce_review_scenario_and_preserve_original(self):
        before = copy.deepcopy(self.case)
        strict = calc.calculate_legal(self.case, self.packet['inputs'])
        result = calc.calculate_provisional(self.case)
        self.assertEqual(strict['schedule'], [])
        self.assertEqual(result['status'], 'provisional')
        self.assertEqual(result['summary']['months'], 36)
        self.assertEqual(result['summary']['total_creditor_payment'], 45412452)
        self.assertFalse(result['submission_ready'])
        self.assertIsNone(result['approval'])
        self.assertTrue(result['pending_conditions'])
        self.assertTrue(all(a['status'] == 'review_required' for a in result['assumptions']))
        self.assertTrue(calc.validate_provisional(self.case, result))
        self.assertEqual(self.case, before)
        self.assertIsNone(self.packet['inputs']['assets'][0]['exempt_deduction'])

    def test_explicit_scenario_preserves_period_and_deductions_but_never_salary_override(self):
        supplied = {'months': 60, 'period_exception': 'liquidation',
            'income': {'monthly_amount': 99000000},
            'assets': [{'id': self.inputs['assets'][0]['id'], 'exempt_deduction': 1000000}]}
        result = calc.calculate_provisional(self.case, supplied)
        self.assertEqual(result['inputs']['months'], 60)
        self.assertEqual(result['inputs']['income']['monthly_amount'], 2800000)
        self.assertEqual(result['summary']['liquidation_value'], 11000000)
        self.assertIn('income.monthly_amount', result['ignored_input_fields'])
        self.assertIn('SOURCE_INPUT_DIFFERENCE', {r['code'] for r in result['pending_conditions']})
        self.assertTrue(calc.validate_provisional(self.case, result))

    def test_missing_real_income_household_or_mapping_conflict_never_creates_usable_plan(self):
        original = copy.deepcopy(self.packet)
        for change in ('income', 'household', 'mapping'):
            self.packet = copy.deepcopy(original)
            if change == 'income':
                self.packet['inputs']['income']['monthly_amount'] = None
            elif change == 'household':
                self.packet['form_values'] = {}
            else:
                self.packet['errors'] = [{'code': 'SOURCE_CONFLICT', 'reason': 'synthetic conflict'}]
            result = calc.calculate_provisional(self.case)
            self.assertEqual(result['status'], 'blocked')
            self.assertFalse(calc.validate_provisional(self.case, result))

    def test_other_court_does_not_inherit_seoul_living_cost(self):
        self.case['court_id'] = 'CT02'
        result = calc.calculate_provisional(self.case)
        self.assertEqual(result['status'], 'blocked')
        self.assertIsNone(result['inputs']['living_cost_mode'])

    def test_tampered_summary_assumption_or_source_revision_is_rejected(self):
        result = calc.calculate_provisional(self.case)
        for field in ('summary', 'assumptions', 'approval', 'status', 'input_revision'):
            changed = copy.deepcopy(result)
            if field == 'summary':
                changed[field]['monthly_deposit'] += 1
            elif field == 'assumptions':
                changed[field] = []
            else:
                changed[field] = {'approved': True} if field == 'approval' else 'approved' if field == 'status' else 99
            self.assertFalse(calc.validate_provisional(self.case, changed), field)
        self.packet['source_signature'] = 'changed-source'
        self.assertFalse(calc.validate_provisional(self.case, result))

    def test_unknown_objection_does_not_become_no_objection_and_paid_costs(self):
        result = calc.calculate_provisional(self.case)
        self.assertTrue(result['summary']['minimum_test_applied'])
        self.assertFalse(result['inputs']['preapproval_costs_paid'])
        codes = {r['code'] for r in result['pending_conditions']}
        self.assertIn('OBJECTORS_REQUIRED', codes)
        self.assertIn('PREAPPROVAL_COSTS_UNPAID', codes)

    def test_repeat_preparation_reuses_unique_id_and_preserves_history(self):
        result = calc.calculate_provisional(self.case)
        first = calc.record_provisional(self.case, result, created_at='synthetic-first', analysis_calculation_id='strict-one')
        second = calc.record_provisional(self.case, result, created_at='synthetic-second', analysis_calculation_id='strict-two')
        self.assertEqual(first['id'], second['id'])
        self.assertEqual(len(self.case['legal_calculations']), 1)
        self.assertEqual(len(second['preparation_history']), 2)
        self.assertTrue(calc.validate_provisional(self.case, second))
        explicit = calc.calculate_provisional(self.case, {'months': 36})
        self.assertNotEqual(result['id'], explicit['id'])
        self.assertNotEqual(result['assumptions'][0]['basis'], explicit['assumptions'][0]['basis'])


class ProvisionalDraftHandoffTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_draft_docx_html_and_seven_pdfs_keep_calculation_values(self):
        from docx import Document
        from apps.api import automation, preliminary_drafting, store
        from tests.test_evidence_mapping_pipeline import fixture
        case = fixture()
        with tempfile.TemporaryDirectory() as folder, patch.object(store, 'DATA_DIR', Path(folder)), \
                patch.object(automation, '_retrieve', return_value=[]), \
                patch.object(automation, 'approved_examples', return_value=[]), \
                patch('apps.api.grounded_drafting.compose', new=AsyncMock()) as writer, \
                patch('apps.api.verification.run_local_verification_batched', new=AsyncMock()) as model:
            draft = await preliminary_drafting.create(case, [], {'analysis'},
                local_failure={'status': 'unavailable', 'error': {'code': 'SYNTHETIC_NO_MODEL'}})
            projection = next(c for c in case['legal_calculations'] if c.get('status') == 'provisional')
            section = next(s for s in draft['sections'] if s['id'] == 'repayment_projection')
            values = {f['key']: f['value'] for f in section['fields']}
            self.assertEqual(values['months'], 36)
            self.assertEqual(values['monthly_deposit'], projection['summary']['monthly_deposit'])
            self.assertNotIn(None, values.values())
            self.assertNotIn('seoul_median_60', section['content'])
            self.assertIn('이의가 있는 경우의 요건도 계산', section['content'])
            directory = Path(folder) / 'generated' / case['id']
            html = (directory / (draft['id'] + '.html')).read_text(encoding='utf-8')
            docx = Document(io.BytesIO((directory / (draft['id'] + '.docx')).read_bytes()))
            document_text = '\n'.join(p.text for p in docx.paragraphs) + '\n' + '\n'.join(
                c.text for table in docx.tables for row in table.rows for c in row.cells)
            for text in (html, document_text):
                self.assertIn(str(projection['summary']['monthly_deposit']), text)
                self.assertIn('변제계획 검토안', text)
                self.assertNotIn('seoul_median_60', text)
            self.assertEqual(len(case['court_documents']), 7)
            self.assertTrue(all(d['calculation_id'] == projection['id'] for d in case['court_documents']))
            self.assertTrue(all(not d['submission_ready'] for d in case['court_documents']))
            self.assertFalse(draft['submission_ready'])
            self.assertIsNone(draft['calculation_id'])
            self.assertFalse(draft['ai_review']['passed'])
            writer.assert_not_awaited()
            model.assert_not_awaited()


class ProvisionalApiTests(unittest.TestCase):
    def test_manual_calculation_keeps_strict_result_and_downloads_labelled_projection(self):
        from fastapi import FastAPI
        from fastapi.responses import JSONResponse
        from fastapi.testclient import TestClient
        from apps.api import domain, evidence_mapping, extended_routes, store
        from tests.test_evidence_mapping_pipeline import fixture
        case = fixture()
        case['matter_number'] = 'TEST-2026-001'
        user = {'id': 'synthetic-lawyer', 'role': 'lawyer', 'name': '검증담당자'}
        app = FastAPI()
        @app.exception_handler(domain.DomainError)
        async def domain_error(request, error):
            return JSONResponse(status_code=409, content={'code': error.code})
        def change(case_id, actor, version, action, apply):
            apply(case)
            return case
        extended_routes.attach(app, lambda: user, lambda case_id, actor: case, change)
        with TestClient(app) as client, patch.object(store, 'access'):
            response = client.post(f"/api/cases/{case['id']}/legal-calculations", json={
                'expected_version': case['version'], 'inputs': evidence_mapping.build(case)['inputs']})
            self.assertEqual(response.status_code, 200)
            strict, projection = response.json()['legal_calculations'][-2:]
            self.assertEqual(strict['status'], 'blocked')
            self.assertEqual(projection['status'], 'provisional')
            self.assertEqual(projection['analysis_calculation_id'], strict['id'])
            download = client.get(f"/api/cases/{case['id']}/legal-calculations/{projection['id']}/download")
            self.assertEqual(download.status_code, 200)
            text = download.content.decode('utf-8-sig')
            self.assertIn('검토용 가정 계산', text)
            self.assertIn('서울 기준 중위소득의 60%', text)
            self.assertNotIn('seoul_median_60', text)
            self.assertNotIn('입력 해시', text)
            self.assertIn('TEST-2026-001_', download.headers['content-disposition'])
            approval = client.post(f"/api/cases/{case['id']}/legal-calculations/{projection['id']}/approve", json={
                'expected_version': case['version'], 'reason': '가정만으로 승인 불가'})
            self.assertGreaterEqual(approval.status_code, 400)
            self.assertIsNone(case['legal_calculations'][-1]['approval'])


if __name__ == '__main__':
    unittest.main()
