"""Create reproducible, explicitly fictional evidence for local end-to-end testing."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
import sys
import zipfile

import pymupdf
from PIL import Image
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'examples' / 'synthetic_case'
FONT = Path('C:/Windows/Fonts/malgun.ttf')
WARNING = '가상자료 · 효력없음 · 로컬 테스트 전용'


def pdf_document(path, title, pages, *, scanned=False, save_png=False):
    doc = pymupdf.open()
    for number, lines in enumerate(pages, 1):
        page = doc.new_page(width=595, height=842)
        page.insert_font(fontname='Korean', fontfile=str(FONT))
        page.draw_rect((35, 30, 560, 74), color=(.7,.15,.1), fill=(1,.95,.91))
        page.insert_text((48, 56), WARNING, fontsize=12, fontname='Korean', color=(.7,.15,.1))
        page.insert_text((45, 111), title, fontsize=20, fontname='Korean')
        y=150
        for line in lines:
            if isinstance(line, (list, tuple)):
                # Clear, aligned table rows; no imitation official seals or signatures.
                count=len(line)
                for col,value in enumerate(line):
                    x=45+col*505/count
                    page.draw_rect((x,y-17,x+505/count,y+15),color=(.7,.75,.8),width=.5)
                    page.insert_text((x+6,y+3), str(value), fontsize=10, fontname='Korean')
                y+=32
            else:
                page.insert_text((45,y),str(line),fontsize=12,fontname='Korean')
                y+=29
        page.insert_text((45,795),f'{WARNING} / {number}쪽 / 실제 발급기관·계좌·서명 없음',fontsize=9,fontname='Korean',color=(.55,.15,.1))
    if scanned:
        image_doc=pymupdf.open()
        for n,page in enumerate(doc):
            pix=page.get_pixmap(matrix=pymupdf.Matrix(2.5,2.5),alpha=False)
            image=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
            stream=io.BytesIO();image.save(stream,format='JPEG',quality=88)
            raster=image_doc.new_page(width=595,height=842)
            raster.insert_image(raster.rect,stream=stream.getvalue())
            if save_png and n==0:image.save(path.with_suffix('.png'))
        image_doc.save(path,deflate=True)
        image_doc.close()
    else:doc.subset_fonts();doc.save(path,deflate=True)
    doc.close()


def generate():
    docs=OUT/'documents';docs.mkdir(parents=True,exist_ok=True)
    consultation=f'''{WARNING}
상담일: 2026-10-03. 내담자 성명: 김가온. 거주지: 서울특별시 가상구 예시로 100, 테스트동 101호(존재하지 않는 주소).
개인회생 상담을 요청합니다. 예시테크(가상회사) 정규직 급여소득자이며 2024년 1월부터 근무했습니다.
2026년 7월, 8월, 9월의 급여는 매월 총지급액 3,100,000원, 공제액 300,000원, 실수령 월소득 2,800,000원입니다.
채무총액은 78,000,000원이고 원금 75,000,000원, 이자 3,000,000원입니다. 모두 무담보 채무입니다.
예시은행A 원금 35,000,000원 이자 1,400,000원, 예시카드B 원금 25,000,000원 이자 1,000,000원, 예시저축C 원금 15,000,000원 이자 600,000원입니다.
실제 월지출은 월세 550,000원, 식비 450,000원, 교통비 150,000원, 통신비 80,000원, 공과금 120,000원, 기타 150,000원으로 합계 1,500,000원입니다.
임차보증금 10,000,000원, 예금 1,200,000원, 보험해약환급금 800,000원으로 재산 합계 12,000,000원입니다. 부동산과 자동차는 없습니다.
미혼이며 1인 가구이고 부양가족은 없습니다. 연체가 시작되어 정기소득으로 분할변제하려고 합니다.
2026-08-10 가족 김예시가 예시카드B에 2,000,000원을 대신 상환했습니다. 증여가 아니라 가족에게 갚아야 하는 차용금입니다.
주의: 위 78,000,000원에는 가족 채무 2,000,000원이 포함되지 않았다는 상담상 누락이 있습니다. 가족 지원의 법률관계와 채권자 추가 여부를 반드시 확인해야 합니다.
최근 1년 신규차입, 재산처분, 세금체납, 과거 회생·파산·면책은 없다고 진술합니다. 이 진술은 제출자료와 대조해야 합니다.
주민등록번호는 TEST-ID-001로 표시하고 실제 번호는 사용하지 않습니다. 연락처와 서명은 공란으로 둡니다.
'''
    (OUT/'00_상담회의록_가상자료.txt').write_text(consultation,encoding='utf-8')
    docx=Document();docx.add_heading('개인회생 상담회의록 · 가상자료',0)
    for line in consultation.splitlines():docx.add_paragraph(line)
    docx.save(OUT/'00_상담회의록_가상자료.docx')
    salary=[]
    for month in (7,8,9):
        salary.append(['성명: 김가온',f'귀속연월: 2026년 {month:02d}월 / 지급일: 2026-{month:02d}-25','회사명: 예시테크(가상회사) / 고용형태: 정규직',
          ['지급항목','금액(원)'],['기본급','2,900,000'],['식대','200,000'],['총지급액','3,100,000'],
          ['공제 합계','300,000'],['차인지급액(실수령액)','2,800,000'],
          '월소득: 2,800,000원','지급계좌: TEST-ACCOUNT-001 (사용 불가능한 테스트 식별자)','공제내역은 테스트 합계이며 실제 보험료·세액 산식이 아닙니다.','발급: 가상회사 테스트 담당 / 날인 없음'])
    pdf_document(docs/'01_급여명세서_가상자료.pdf','급여명세서 · 스캔 테스트',salary,scanned=True,save_png=True)
    transactions=[]; balance=1200000
    for month in (7,8,9):
        for day,desc,deposit,withdrawal in [(25,'급여 예시테크',2800000,0),(26,'월세',0,550000),(27,'생활비 합계',0,950000),(28,'채무 상환 합계',0,1300000)]:
            balance+=deposit-withdrawal
            transactions.append({'date':f'2026-{month:02d}-{day:02d}','description':desc,'deposit':deposit,'withdrawal':withdrawal,'balance':balance})
    with (docs/'02_거래내역_가상자료.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=['date','description','deposit','withdrawal','balance']);writer.writeheader();writer.writerows(transactions)
    pages=[]
    for month in (7,8,9):
        rows=[x for x in transactions if x['date'].startswith(f'2026-{month:02d}')]
        pages.append(['예금주: 김가온 / 금융기관: 예시은행A(가상기관)','계좌번호: TEST-ACCOUNT-001',f'조회기간: 2026-{month:02d}-01 ~ 2026-{month:02d}-{31 if month in (7,8) else 30}',
          '기초잔액: 1,200,000원',['일자','입금','출금','잔액']]+[[r['date'],f"{r['deposit']:,}",f"{r['withdrawal']:,}",f"{r['balance']:,}"] for r in rows]+['기말잔액: 1,200,000원','생활비와 상환 합계는 가상 테스트 분류입니다.'])
    pdf_document(docs/'02_계좌거래내역_가상자료.pdf','입출금 거래내역',pages)
    creditors=[{'name':'예시은행A','principal':35000000,'interest':1400000,'secured':False}, {'name':'예시카드B','principal':25000000,'interest':1000000,'secured':False},{'name':'예시저축C','principal':15000000,'interest':600000,'secured':False}]
    pdf_document(docs/'03_부채증명서_가상자료.pdf','부채증명서 · 테스트 양식',[
      ['채무자: 김가온','기준일: 2026-09-30',f"채권자: {c['name']} (가상기관)",['구분','금액(원)'],['대출원금',f"{c['principal']:,}"],['이자',f"{c['interest']:,}"],['합계',f"{c['principal']+c['interest']:,}"],'담보: 없음 / 보증인: 없음','계약번호: TEST-LOAN-00'+str(i+1),'본 자료는 실제 금융기관의 발급서류가 아닙니다.'] for i,c in enumerate(creditors)])
    pdf_document(docs/'04_주민등록_가족관계_가상자료.pdf','주소·가족관계 확인용 가상자료',[
      ['성명: 김가온','식별번호: TEST-ID-001','주민등록상 주소: 서울특별시 가상구 예시로 100','상세주소: 테스트동 101호 (존재하지 않는 주소)','세대원 수: 1명 / 세대주: 김가온','전입일: 2025-01-01','혼인: 미혼 / 부양가족: 없음','테스트에 필요한 항목만 구성한 예시입니다.','주민등록등본·가족관계증명서 정식 양식이나 발급문서가 아닙니다.']])
    pdf_document(docs/'05_임대차계약서_가상자료.pdf','주거용 임대차계약서 · 가상자료',[
      ['임차인: 김가온 / 임대인: 박가상','소재지: 서울특별시 가상구 예시로 100 테스트동 101호', '임대차기간: 2025-01-01 ~ 2026-12-31',['항목','계약금액'],['임차보증금','10,000,000원'],['월 차임','550,000원'],'차임 지급일: 매월 26일','확정일자·전입의 법적 효력은 테스트 데이터에 포함하지 않습니다.','임대인 서명: 공란 / 임차인 서명: 공란']])
    pdf_document(docs/'06_보험해약환급금_가상자료.pdf','보험해약환급금 예상액',[
      ['계약자: 김가온 / 보험사: 예시보험(가상)','보험상품: 테스트 보장보험 / 증권번호: TEST-POLICY-001','기준일: 2026-09-30','해약환급금: 800,000원','보험계약대출: 0원','합계 순환급예상액: 800,000원','실제 보험상품·증권·발급서류가 아닙니다.']])
    pdf_document(docs/'07_가족대납_가상자료.pdf','가족 대납 사실확인 · 가상자료',[
      ['채무자: 김가온 / 대납자: 김예시','상환일: 2026-08-10 / 수취인: 예시카드B','상환액: 2,000,000원','자금 출처: 가족 김예시의 자기자금','법률관계 진술: 증여가 아니라 상환 의무가 있는 차용','가족 채무 잔액: 2,000,000원','금융기관 3곳의 기준일 잔액 78,000,000원과 별도입니다.','따라서 가족채무가 인정되면 총채무는 80,000,000원입니다.','대납자 서명 없음. 사실 확인 및 채권자 포함 여부 검토 필요.']])
    edge=OUT/'edge_cases';edge.mkdir(exist_ok=True)
    pdf_document(edge/'다른사람_급여명세서_가상자료.pdf','급여명세서 · 다른 사람 검출 테스트',[
      ['성명: 이다른','귀속연월: 2026년 09월','회사명: 예시테크(가상회사)','월소득: 9,900,000원','김가온 사건에 넣으면 인적 동일성 충돌로 격리되어야 합니다.']],scanned=True)
    pdf_document(edge/'기간부족_급여명세서_가상자료.pdf','급여명세서 · 기간부족 테스트',[salary[-1]],scanned=True)
    expected={'synthetic':True,'name':'김가온','court_id':'CT01','as_of':'2026-10-03','monthly_income':2800000,'gross_income':3100000,'monthly_deductions':300000,'living_expenses':1500000,'household_size':1,'assets_total':12000000,'lease_deposit':10000000,'bank_balance':1200000,'insurance_surrender':800000,'bank_debt_total':78000000,'principal_total':75000000,'interest_total':3000000,'family_debt_candidate':2000000,'total_debt_after_family_review':80000000,'creditors':creditors,'transactions':transactions,'expected_rules':['WF01','WF02','WF04','WF05','WF12'],'intended_findings':['가족채무 누락 2000000원 확인','금융기관 채무 78000000원과 가족채무 분리','급여 3개월은 1년 요청에 기간부족','다른사람 문서는 격리'],'submission_order':['00_상담회의록_가상자료.txt','documents/01_급여명세서_가상자료.pdf','documents/02_계좌거래내역_가상자료.pdf','documents/03_부채증명서_가상자료.pdf','documents/04_주민등록_가족관계_가상자료.pdf','documents/05_임대차계약서_가상자료.pdf','documents/06_보험해약환급금_가상자료.pdf','documents/07_가족대납_가상자료.pdf']}
    (OUT/'expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'README.txt').write_text(WARNING+'\n회의록을 먼저 넣고 documents 폴더의 PDF를 차례로 올리세요. 급여 PDF는 텍스트층 없는 이미지형 스캔입니다.\nexpected.json은 정답 비교용입니다. 증빙으로 업로드하지 마세요.\n가족 차용금 200만원 누락을 의도적으로 넣었습니다. 은행부채7800만원→가족 포함8000만원 변경과 초안 갱신을 확인하세요.\nedge_cases는 정상 자료 처리 후 별도 사건에서 테스트하세요.\n정식 정부증명서의 위조본이 아닌 필드 구성에 맞춘 테스트 양식입니다. 실제 금융기관·계좌·주민번호·서명·관인은 없습니다.\n',encoding='utf-8')
    form_fields=OUT/'sample-form-fields.json'
    if form_fields.exists():
        sys.path.insert(0,str(ROOT))
        from apps.api.court_forms import render_pdf
        fields=json.loads(form_fields.read_text(encoding='utf-8'))
        form_case={'id':'synthetic','client_name':'김가온','court_id':'CT01','court_name':'서울회생법원','case_type':'personal_rehabilitation','input_revision':1,'synthetic':True,'consultation':{'notes':consultation}}
        preview_dir=OUT/'court_previews';preview_dir.mkdir(exist_ok=True)
        for template in ['D5100','D5101','D5103','D5105','D5106','D5110','D5115']:
            (preview_dir/(template+'_가상자료_검토초안.pdf')).write_bytes(render_pdf(template,form_case,fields))
    legal_example=ROOT/'examples'/'legal_inputs.json'
    if legal_example.exists():
        inputs=json.loads(legal_example.read_text(encoding='utf-8'))
        inputs['decisions']['debt_reason']='가상 부채증명서 3곳 7800만원만 반영한 비교안. 별도 가족채무200만원의 채권자 추가 검토가 남아 있으므로 최종 확정 금액이 아닙니다.'
        (OUT/'legal_inputs_bank_only.json').write_text(json.dumps(inputs,ensure_ascii=False,indent=2),encoding='utf-8')
        inputs['creditors'].append({'id':'creditor-family','name':'김예시(가상)','kind':'unsecured','principal':2000000,'interest':0,'evidence_ids':['fixture-family']})
        inputs['decisions']['debt_reason']='가상 사례에서 가족 대납이 상환의무 있는 차용으로 확인되었다는 비교안. 가족채무200만원 포함 총채무8000만원. 실제 운영에서는 증빙검토·변호사승인 전 확정하지 않습니다.'
        (OUT/'legal_inputs_family_included.json').write_text(json.dumps(inputs,ensure_ascii=False,indent=2),encoding='utf-8')
    zip_path=ROOT/'examples'/'synthetic_case_bundle.zip'
    with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as archive:
        for file in sorted(OUT.rglob('*')):
            if file.is_file():archive.write(file,file.relative_to(OUT))
    print(json.dumps({'output':str(OUT),'zip':str(zip_path),'files':len(list(OUT.rglob('*'))),'synthetic':True},ensure_ascii=False))


if __name__=='__main__':generate()
